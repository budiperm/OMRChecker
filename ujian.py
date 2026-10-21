"""

 Exam ("ujian") logic for the OMRChecker wizard, without any GUI:

 - template library (templates kept in the app data folder, not in scan folders)
 - exam settings that persist between sessions
 - scoring with one answer key per exam code (kode naskah), manual corrections
 - matching against a participant list (Excel/CSV)
 - item analysis and the Excel export

"""

import csv
import hashlib
import json
import os
import re
import shutil
import time
from copy import deepcopy
from dataclasses import asdict, dataclass, field
from pathlib import Path

EXAM_FILE = "ujian.json"  # exam settings kept in its output folder
ARCHIVE_DIR = "_arsip"  # archived scans (and their results): same as src.constants ARCHIVE_DIRNAME
ALL_CODES = "*"  # key used for every sheet when there is no exam-code column
OPTIONS_ORDER = ["A", "B", "C", "D", "E"]


def app_data_dir():
    base = os.environ.get("APPDATA") or Path.home().joinpath(".config")
    return Path(base, "OMRChecker")


def library_dir():
    return app_data_dir() / "templates"


def safe_name(name):
    name = re.sub(r'[<>:"/\\|?*]+', " ", name).strip().rstrip(".")
    return name or "Template"


def natural_key(text):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(text))]


# ---------------------------------------------------------------------------
# Template library
# ---------------------------------------------------------------------------

BLANK_PREVIEW = "ljk_kosong.png"
BLOCKS_FILE = "blok.json"
META_FILE = "info.json"


@dataclass
class LibraryTemplate:
    path: Path
    name: str
    sections: dict  # part name -> question labels in order ("" = no parts)
    columns: list  # identity columns, e.g. Nama, NIS, KodeSoal
    created: str = ""
    source: str = ""  # file the template was made from

    @property
    def questions(self):
        return [q for labels in self.sections.values() for q in labels]

    @property
    def preview(self):
        p = self.path / BLANK_PREVIEW
        return p if p.exists() else None

    def summary(self):
        cols = ", ".join(self.columns) if self.columns else "-"
        if len(self.sections) > 1 or "" not in self.sections:
            parts = ", ".join(f"{n}: {len(q)}" for n, q in self.sections.items())
            return f"{len(self.questions)} soal ({parts}); kolom: {cols}"
        return f"{len(self.questions)} soal; kolom: {cols}"


def _expand(label):
    m = re.fullmatch(r"(.+?)(\d+)(?:\.\.(\d+))?", label)
    if not m:
        return []
    prefix, a, b = m.group(1), int(m.group(2)), int(m.group(3) or m.group(2))
    return [f"{prefix}{i}" for i in range(a, b + 1)]


def template_sections(template_json, names=None):
    """{part name: labels}. `names` maps label prefix -> part name (from
    bagian.json); without it every q-label belongs to one unnamed part."""
    names = names or {"q": ""}
    found = {}
    for block in template_json.get("fieldBlocks", {}).values():
        for s in block.get("fieldLabels", []):
            for label in _expand(s):
                prefix = label.rstrip("0123456789")
                if prefix in names:
                    found.setdefault(names[prefix], set()).add(label)
    order = list(dict.fromkeys(names.values()))
    return {n: sorted(found[n], key=natural_key) for n in order if n in found}


def template_questions(template_json, names=None):
    return [q for labels in template_sections(template_json, names).values() for q in labels]


def template_columns(template_json):
    return list(template_json.get("customLabels", {}).keys())


def read_library_template(path):
    path = Path(path)
    if not (path / "template.json").exists():
        raise ValueError(f"Bukan folder template: {path}")
    t = json.loads((path / "template.json").read_text(encoding="utf-8"))
    meta = {}
    if (path / META_FILE).exists():
        meta = json.loads((path / META_FILE).read_text(encoding="utf-8"))
    return LibraryTemplate(
        path=path,
        name=meta.get("name", path.name),
        sections=template_sections(t, section_names(path)),
        columns=template_columns(t),
        created=meta.get("created", ""),
        source=meta.get("source_name", ""),
    )


def section_names(folder):
    import template_builder as tb

    p = Path(folder) / tb.SECTIONS_FILE
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return None


def list_templates(lib=None):
    lib = Path(lib or library_dir())
    if not lib.exists():
        return []
    found = []
    for d in sorted(lib.iterdir(), key=lambda p: natural_key(p.name)):
        if (d / "template.json").exists():
            try:
                found.append(read_library_template(d))
            except (OSError, ValueError):
                continue
    return found


def save_to_library(name, sheet, template, blocks, lib=None, replace=None):
    """Store a template made in the editor. Returns its folder.

    `replace`: folder of the template being edited (overwritten in place)."""
    import cv2

    import template_builder as tb

    lib = Path(lib or library_dir())
    if replace:
        folder = Path(replace)
    else:
        folder = lib / safe_name(name)
        n = 2
        while folder.exists():
            folder = lib / f"{safe_name(name)} ({n})"
            n += 1
    folder.mkdir(parents=True, exist_ok=True)
    for stale in ("template.json", "config.json", tb.REFERENCE_FILENAME, tb.ORIENTATION_FILENAME):
        (folder / stale).unlink(missing_ok=True)
    (folder / tb.SECTIONS_FILE).unlink(missing_ok=True)
    tb.save_template_folder(folder, sheet, template, sections=tb.sections_of(blocks))
    ok, buf = cv2.imencode(".png", sheet.page)
    buf.tofile(str(folder / BLANK_PREVIEW))
    (folder / BLOCKS_FILE).write_text(
        json.dumps([asdict(b) for b in blocks], indent=1), encoding="utf-8"
    )
    meta = {
        "name": name,
        "created": time.strftime("%Y-%m-%d %H:%M"),
        "markers": sheet.markers,
        "marker_size": sheet.marker_size,
        "source_name": sheet.source_name,
        "source_blank": sheet.source_blank,
        "page_dims": sheet.page_dims,
    }
    (folder / META_FILE).write_text(json.dumps(meta, indent=1), encoding="utf-8")
    return folder


def load_for_editing(folder):
    """(Sheet, blocks, name) to reopen a library template in the editor."""
    import cv2
    import numpy as np

    import template_builder as tb

    folder = Path(folder)
    meta = json.loads((folder / META_FILE).read_text(encoding="utf-8"))
    page = cv2.imdecode(np.fromfile(str(folder / BLANK_PREVIEW), np.uint8), cv2.IMREAD_COLOR)
    sheet = tb.Sheet(
        page,
        meta["page_dims"],
        markers=meta.get("markers"),
        marker_size=meta.get("marker_size", 0),
        source_name=meta.get("source_name", ""),
        source_blank=meta.get("source_blank", True),
    )
    blocks = [tb.Block(**b) for b in json.loads((folder / BLOCKS_FILE).read_text(encoding="utf-8"))]
    return sheet, blocks, meta.get("name", folder.name)


def section_box(folder, part):
    """(x0, y0, x1, y1) around a part's questions on the blank preview, or
    None for templates without block data (imported ones)."""
    folder = Path(folder)
    if not (folder / BLOCKS_FILE).exists():
        return None
    blocks = json.loads((folder / BLOCKS_FILE).read_text(encoding="utf-8"))
    mine = [b for b in blocks if b.get("kind") == "mcq" and b.get("section", "") == part]
    if not mine:
        return None
    x0 = min(b["x0"] for b in mine)
    y0 = min(b["y0"] for b in mine)
    x1 = max(b["x0"] + b["gap_x"] * (b["cols"] - 1) + b["bw"] for b in mine)
    y1 = max(b["y0"] + b["gap_y"] * (b["rows"] - 1) + b["bh"] for b in mine)
    return x0, y0, x1, y1


def section_thumbnail(folder, part, margin=12):
    """Crop of the blank sheet showing where a part is (BGR) or None."""
    import cv2
    import numpy as np

    box = section_box(folder, part)
    preview = Path(folder) / BLANK_PREVIEW
    if box is None or not preview.exists():
        return None
    img = cv2.imdecode(np.fromfile(str(preview), np.uint8), cv2.IMREAD_COLOR)
    h, w = img.shape[:2]
    x0, y0, x1, y1 = (int(round(v)) for v in box)
    return img[max(0, y0 - margin) : min(h, y1 + margin), max(0, x0 - margin) : min(w, x1 + margin)].copy()


def rename_section(folder, old, new):
    """Rename a part of a library template (labels change with the name).
    Returns the {old label: new label} mapping for anything keyed by label."""
    import template_builder as tb

    folder = Path(folder)
    if not can_edit(folder):
        raise ValueError("Template hasil impor tidak bisa diganti nama bagiannya.")
    new = new.strip()
    sheet, blocks, name = load_for_editing(folder)
    if any(b.kind == tb.KIND_MCQ and b.section == new for b in blocks) and new != old:
        raise ValueError(f"Bagian '{new}' sudah ada.")
    before = tb.question_labels(blocks)
    for b in blocks:
        if b.kind == tb.KIND_MCQ and b.section == old:
            b.section = new
    template, problems = tb.build_template(blocks, sheet)
    if problems:
        raise ValueError("; ".join(problems))
    after = tb.question_labels(blocks)
    old_prefix, new_prefix = tb.section_prefix(old), tb.section_prefix(new)
    mapping = {
        label: new_prefix + label[len(old_prefix):]
        for label in before
        if label.startswith(old_prefix) and new_prefix + label[len(old_prefix):] in after
    }
    save_to_library(name, sheet, template, blocks, folder.parent, replace=folder)
    return mapping


def rename_section_in_exam(exam, old, new, mapping):
    """Move keys, weights, the in-order rule and corrections of a renamed part."""
    for code, keys in list(exam.keys.items()):
        if isinstance(keys, dict) and old in keys:
            keys = dict(keys)
            keys[new] = keys.pop(old)
            exam.keys[code] = keys
    if old in exam.weights:
        exam.weights[new] = exam.weights.pop(old)
    if isinstance(exam.in_order, list):
        exam.in_order = [new if p == old else p for p in exam.in_order]
    else:
        exam.in_order = {c: [new if p == old else p for p in parts] for c, parts in exam.in_order.items()}
    for fid, fixes in exam.corrections.items():
        exam.corrections[fid] = {mapping.get(k, k): v for k, v in fixes.items()}


def upgrade_template(folder):
    """Re-save an editable library template made by an older version so it
    gets newer safeguards (e.g. turning upside-down sheets). Returns True
    when something changed."""
    import template_builder as tb

    folder = Path(folder)
    if not can_edit(folder):
        return False
    t = json.loads((folder / "template.json").read_text(encoding="utf-8"))
    pre = t.get("preProcessors", [{}])[0]
    if pre.get("name") != "AlignOnMarkers" or "orientationReference" in pre.get("options", {}):
        return False
    sheet, blocks, name = load_for_editing(folder)
    template, problems = tb.build_template(blocks, sheet)
    if problems:
        return False
    save_to_library(name, sheet, template, blocks, folder.parent, replace=folder)
    return True


def can_edit(folder):
    folder = Path(folder)
    return all((folder / f).exists() for f in (META_FILE, BLOCKS_FILE, BLANK_PREVIEW))


def import_template_folder(src, name, lib=None):
    """Bring a template folder from the advanced mode into the library."""
    src, lib = Path(src), Path(lib or library_dir())
    if not (src / "template.json").exists():
        raise ValueError("Folder ini tidak berisi template.json")
    folder = lib / safe_name(name)
    n = 2
    while folder.exists():
        folder = lib / f"{safe_name(name)} ({n})"
        n += 1
    folder.mkdir(parents=True)
    for f in ("template.json", "config.json", "ljk_referensi.png", "ljk_arah.png", "bagian.json"):
        if (src / f).exists():
            shutil.copy2(src / f, folder / f)
    if (src / "ljk_referensi.png").exists():
        shutil.copy2(src / "ljk_referensi.png", folder / BLANK_PREVIEW)
    (folder / META_FILE).write_text(
        json.dumps({"name": name, "created": time.strftime("%Y-%m-%d %H:%M")}), encoding="utf-8"
    )
    return folder


TEMPLATE_EXT = ".ljk"  # one-file template: a zip of the library folder


def export_template_file(folder, path):
    """Pack a library template into one .ljk file to use on another computer."""
    import zipfile

    folder, path = Path(folder), Path(path)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(folder.iterdir()):
            if f.is_file():
                z.write(f, f.name)
    return path


def import_template_file(path, lib=None):
    """Unpack a .ljk file into the library; returns the new folder. A plain
    template.json (advanced mode) brings in the folder it is in."""
    import zipfile

    path, lib = Path(path), Path(lib or library_dir())
    if path.name == "template.json":
        return import_template_folder(path.parent, path.parent.name, lib)
    try:
        z = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as e:
        raise ValueError(f"Bukan file template LJK ({TEMPLATE_EXT}): {path.name}") from e
    with z:
        names = [n for n in z.namelist() if not n.endswith("/")]
        if "template.json" not in names:
            raise ValueError(f"File {path.name} tidak berisi template.json")
        if any(Path(n).name != n or n.startswith(".") for n in names):  # no folders, no ../
            raise ValueError(f"File {path.name} berisi nama file yang tidak wajar, tidak diimpor.")
        meta = json.loads(z.read(META_FILE).decode("utf-8")) if META_FILE in names else {}
        name = meta.get("name") or path.stem
        folder = lib / safe_name(name)
        n = 2
        while folder.exists():
            folder = lib / f"{safe_name(name)} ({n})"
            n += 1
        folder.mkdir(parents=True)
        for n_ in names:
            (folder / n_).write_bytes(z.read(n_))
    try:
        read_library_template(folder)
    except (OSError, ValueError) as e:
        shutil.rmtree(folder, ignore_errors=True)
        raise ValueError(f"Template di {path.name} rusak: {e}") from e
    return folder


def delete_template(folder):
    shutil.rmtree(folder)


# ---------------------------------------------------------------------------
# Exam settings
# ---------------------------------------------------------------------------


@dataclass
class Exam:
    name: str = ""
    template: str = ""  # library folder
    code_column: str = ""  # column holding the exam code; "" = one key for all
    # code -> {part: "ABCDE..."} (ALL_CODES for all sheets; part "" when the
    # template has no parts). A plain string is accepted as {"": string}.
    keys: dict = field(default_factory=dict)
    correct: float = 1
    incorrect: float = 0
    unmarked: float = 0
    scale_100: bool = True
    kkm: float = 0  # 0 = no pass mark
    scans_dir: str = ""
    output_dir: str = ""
    students_file: str = ""
    students_cols: dict = field(default_factory=dict)  # nis / nama / kelas -> header
    corrections: dict = field(default_factory=dict)  # file_id -> {column: value}
    # part -> "1-10:2, 11-20:1": points for a correct answer per question range;
    # questions not listed use `correct`
    weights: dict = field(default_factory=dict)
    # parts that must be answered in order: after the first blank question,
    # the answers that follow are not counted (scored as blank). Per exam
    # code: {code: [part, ...]}; a plain list (v1.3.x) holds for every code
    # (in_order_for).
    in_order: dict = field(default_factory=dict)
    # template folder name -> its KEY_FIELDS and corrections, kept while another
    # template is chosen so switching back brings them back (switch_template)
    template_keys: dict = field(default_factory=dict)
    rekap_order: str = "nama"  # RANK_ORDERS key: row order of the Excel rekap
    exported_file: str = ""  # last Excel rekap written
    exported_state: str = ""  # rekap_state() at that moment

    @staticmethod
    def session_path():
        return app_data_dir() / "ujian_terakhir.json"

    def save(self, path=None):
        """Save as the current session, plus a copy in the output folder (when
        it exists) so this exam can be opened again later (open_exam)."""
        path = Path(path or self.session_path())
        path.parent.mkdir(parents=True, exist_ok=True)
        text = json.dumps(asdict(self), indent=1, ensure_ascii=False)
        path.write_text(text, encoding="utf-8")
        out = Path(self.output_dir) if self.output_dir else None
        if out is not None and out.is_dir():
            try:
                (out / EXAM_FILE).write_text(text, encoding="utf-8")
            except OSError:  # read-only or vanished folder: the session copy is enough
                pass

    @classmethod
    def load(cls, path=None):
        path = Path(path or cls.session_path())
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return cls()
        known = {k: v for k, v in data.items() if k in cls.__dataclass_fields__}
        return cls(**known)


KEY_FIELDS = ("code_column", "keys", "weights", "in_order")  # what belongs to one template
# also put aside per template: corrections name that template's question labels
SWITCHED_FIELDS = KEY_FIELDS + ("corrections",)


def switch_template(exam, folder):
    """Choose another template: put the keys and corrections of the current
    one aside and take the ones saved for `folder` (empty when it never had
    any), so looking at another template and coming back loses nothing."""
    folder = str(folder)
    if folder == exam.template:
        return
    if exam.template:  # no template yet: what was typed belongs to the first one chosen
        exam.template_keys[Path(exam.template).name] = {f: deepcopy(getattr(exam, f)) for f in SWITCHED_FIELDS}
        saved = exam.template_keys.pop(Path(folder).name, {})
        blank = Exam()
        for f in SWITCHED_FIELDS:
            setattr(exam, f, deepcopy(saved.get(f, getattr(blank, f))))
    else:
        exam.corrections = {}
    exam.template = folder


def clear_corrections(exam):
    """Forget every manual correction, also those put aside with another
    template: the sheets they were made on are gone (new scans or results)."""
    exam.corrections = {}
    for saved in exam.template_keys.values():
        saved.pop("corrections", None)


PREFS_FILE = "preferensi.json"  # app preferences, next to the session file
PREF_SCANS = "ingat_folder_scan"  # keep the scan folder for the next exam
PREF_OUTPUT = "ingat_folder_hasil"  # keep the output folder for the next exam
PREF_DASHBOARD = "tampilan_baru"  # open the one-screen look (dashboard.py) instead of the steps
PREF_SENSITIVITY = "kepekaan_baca"  # 0 = Normal ... len(SENSITIVITY_LEVELS) - 1 = Paling peka

# How much darker than the other options of a question a bubble must be to
# count as marked: OMRChecker's MIN_JUMP / CONFIDENT_SURPLUS (defaults 25 / 5).
# Lower reads lightly shaded bubbles and pen crosses, but also smudges.
SENSITIVITY_LEVELS = [
    ("Normal", 25, 5),
    ("Agak peka", 21, 4),
    ("Peka", 17, 4),
    ("Lebih peka", 13, 3),
    ("Paling peka", 10, 2),
]


def sensitivity_level(value):
    try:
        return min(max(int(value), 0), len(SENSITIVITY_LEVELS) - 1)
    except (TypeError, ValueError):
        return 0


def threshold_params(level):
    """config.json threshold_params for a sensitivity level; {} for Normal
    (the template's own config, if any, then applies)."""
    level = sensitivity_level(level)
    if not level:
        return {}
    _, jump, surplus = SENSITIVITY_LEVELS[level]
    return {"MIN_JUMP": jump, "CONFIDENT_SURPLUS": surplus}


def prefs_path(session=None):
    return Path(session or Exam.session_path()).with_name(PREFS_FILE)


def load_prefs(session=None):
    """App preferences; all off unless switched on."""
    prefs = {PREF_SCANS: False, PREF_OUTPUT: False, PREF_DASHBOARD: False, PREF_SENSITIVITY: 0}
    try:
        data = json.loads(prefs_path(session).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return prefs
    if isinstance(data, dict):
        prefs.update({k: bool(data[k]) for k in prefs if k in data and k != PREF_SENSITIVITY})
        prefs[PREF_SENSITIVITY] = sensitivity_level(data.get(PREF_SENSITIVITY, 0))
    return prefs


def save_prefs(prefs, session=None):
    path = prefs_path(session)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(prefs, indent=1), encoding="utf-8")


def forget_folders(exam, prefs):
    """Clear the scan / output folder unless the user asked to remember it,
    so the next exam does not start in a folder that already holds sheets."""
    if not prefs.get(PREF_SCANS):
        exam.scans_dir = ""
    if not prefs.get(PREF_OUTPUT):
        exam.output_dir = ""
        # the rekap exported last time belongs to that output folder: the next
        # exam is "not exported yet" and must not offer to overwrite that file
        exam.exported_file = exam.exported_state = ""
    clear_corrections(exam)


def open_exam(folder, lib=None):
    """The exam saved in an output folder. Raises ValueError with a message
    for the user when the folder holds no exam or its template is gone."""
    folder = Path(folder)
    path = folder / EXAM_FILE
    if not path.is_file():
        raise ValueError(
            "Folder ini bukan folder hasil ujian (tidak ada file ujian.json).\n"
            "Pilih folder hasil yang dipakai di Langkah 4."
        )
    try:
        json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        raise ValueError(f"File ujian.json tidak bisa dibaca:\n{e}") from e
    exam = Exam.load(path)
    exam.output_dir = str(folder)  # the folder may have been moved or copied
    template = Path(exam.template) if exam.template else None
    if template is None or not (template / "template.json").is_file():
        # another computer or user: look for a template of the same name in this library
        same = Path(lib or library_dir()) / (template.name if template else "")
        if template is not None and (same / "template.json").is_file():
            exam.template = str(same)
        else:
            name = template.name if template else "?"
            raise ValueError(f"Template '{name}' untuk ujian ini tidak ada di pustaka template. Impor atau buat dulu template itu.")
    return exam


def rekap_state(exam):
    """Fingerprint of everything the Excel rekap depends on: the result files
    and the exam settings (keys, scoring, corrections, ...). "" when there are
    no results yet. A different value after an export means the rekap is out
    of date (sheets corrected, re-keyed or processed again)."""
    out = Path(exam.output_dir) if exam.output_dir else None
    if out is None or not out.is_dir():
        return ""
    files = sorted((str(p.relative_to(out)), p.stat().st_mtime_ns)
                   for pattern in ("Results/Results_*.csv", "Manual/ErrorFiles.csv") for p in out.rglob(pattern))  # fmt: skip
    if not files:
        return ""
    skip = ("exported_file", "exported_state", "template_keys")  # keys of other templates: not in this rekap
    data = {k: v for k, v in asdict(exam).items() if k not in skip}
    data["files"] = files
    return hashlib.sha1(json.dumps(data, sort_keys=True, default=str).encode("utf-8")).hexdigest()


def parse_weights(text, count=None):
    """'1-10:2, 11-20:1; 25:3' -> ({question number: points}, problems).
    Decimal comma allowed (1,5)."""
    weights, problems = {}, []
    for item in re.split(r"[;\n]|,(?=\s*\d+\s*(?:-\s*\d+\s*)?[:=])", str(text or "")):
        item = item.strip()
        if not item:
            continue
        m = re.fullmatch(r"(\d+)\s*(?:-\s*(\d+))?\s*[:=]\s*(-?\d+(?:[.,]\d+)?)", item)
        if not m:
            problems.append(f"'{item}' tidak dikenali (contoh: 1-10:2)")
            continue
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        points = float(m.group(3).replace(",", "."))
        if a < 1 or b < a:
            problems.append(f"rentang '{item}' tidak valid")
            continue
        if count is not None and b > count:
            problems.append(f"'{item}': soal hanya sampai {count}")
            continue
        for n in range(a, b + 1):
            if n in weights:
                problems.append(f"soal {n} punya dua bobot")
            weights[n] = points
    return weights, problems


def parse_key(text):
    """'1A 2B, cDE' -> 'ABCDE'; also returns characters that are not A-E."""
    letters, invalid = [], []
    for ch in str(text).upper():
        if ch.isspace() or ch.isdigit() or ch in ".,;:-_()/":
            continue
        (letters if ch in OPTIONS_ORDER else invalid).append(ch)
    return "".join(letters), invalid


# ---------------------------------------------------------------------------
# Answer key files: the keys of step 2 saved to use again with the same LJK
# template when the questions (and so the keys) change
# ---------------------------------------------------------------------------

KEYS_EXT = ".kunci"
KEYS_FORMAT = "omrchecker-kunci"


def export_keys_file(exam, sections, path, template_name=""):
    """Save the scored keys of `exam` (per exam code when it reads a code
    column), with weights, the in-order parts and the scoring, as JSON."""
    codes = active_codes(exam)
    data = {
        "format": KEYS_FORMAT,
        "version": 1,
        "template": template_name,
        "sections": {p: len(labels) for p, labels in as_sections(sections).items()},
        "code_column": exam.code_column,
        "keys": {c: keys_for(exam, c) or {} for c in codes},
        "in_order": {c: list(in_order_for(exam, c)) for c in codes},
        "weights": dict(exam.weights),
        "scoring": {"correct": exam.correct, "incorrect": exam.incorrect, "unmarked": exam.unmarked,
                    "scale_100": exam.scale_100, "kkm": exam.kkm},
    }  # fmt: skip
    path = Path(path)
    path.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    return path


def import_keys_file(path, exam, sections, columns=()):
    """Take the keys of a file saved by export_keys_file into `exam` for the
    template whose parts are `sections` ({part: labels}) and whose text
    columns are `columns`. Replaces the keys of the mode the file was saved
    in (the other mode's keys are kept, see active_codes). Returns warnings
    for what did not fit the template; raises ValueError when the file is
    not an answer key file."""
    path = Path(path)
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as e:
        raise ValueError(f"Bukan file kunci jawaban ({KEYS_EXT}): {path.name}") from e
    if not isinstance(data, dict) or data.get("format") != KEYS_FORMAT or not isinstance(data.get("keys"), dict):
        raise ValueError(f"Bukan file kunci jawaban ({KEYS_EXT}): {path.name}")
    sections = as_sections(sections)
    warnings = []
    file_parts = {p for keys in data["keys"].values() if isinstance(keys, dict) for p in keys}
    file_parts |= set(data.get("sections") or {})
    rename = {}
    if len(file_parts) == 1 and len(sections) == 1 and not file_parts & set(sections):
        rename = {next(iter(file_parts)): next(iter(sections))}  # one part, named differently

    def part_of(p):
        return rename.get(p, p)

    def label(p):
        return f"Bagian {p}" if p else "Kunci"

    code_column = str(data.get("code_column") or "")
    if code_column and code_column not in columns:
        guess = next((c for c in columns if "kode" in c.lower()), columns[0] if columns else "")
        warnings.append(f"Kolom kode naskah '{code_column}' tidak ada di template ini"
                        + (f", dipakai kolom '{guess}'. Periksa pilihan kolomnya." if guess else "."))  # fmt: skip
        code_column = guess or code_column
    keys, in_order, extra, long = {}, {}, set(), set()
    file_order = data.get("in_order") if isinstance(data.get("in_order"), dict) else {}
    for orig, parts in data["keys"].items():
        code = str(orig)
        if not code_column:
            code = ALL_CODES
        elif code == ALL_CODES:
            continue
        if isinstance(parts, str):
            parts = {"": parts}
        if not isinstance(parts, dict):
            continue
        keys[code] = {}
        for p, key in parts.items():
            key, _ = parse_key(key)
            if part_of(p) not in sections:
                if key:
                    extra.add(p)
                continue
            keys[code][part_of(p)] = key
            if len(key) > len(sections[part_of(p)]):
                long.add((part_of(p), len(key)))
        order = file_order.get(orig) or []
        in_order[code] = [part_of(p) for p in order if part_of(p) in sections]
    if not any(k for parts in keys.values() for k in parts.values()):
        raise ValueError(f"File {path.name} tidak berisi kunci untuk bagian template ini.")
    warnings += [f"{label(p)} tidak ada di template ini, kuncinya tidak dipakai." for p in sorted(extra)]
    missing = [p for p in sections if not any(parts.get(p) for parts in keys.values())]
    warnings += [f"{label(p)} tidak ada di file ini, kuncinya masih kosong." for p in missing]
    warnings += [f"{label(p)}: kunci berisi {n} jawaban, padahal hanya {len(sections[p])} soal."
                 for p, n in sorted(long)]  # fmt: skip

    if code_column:  # replace the keys of this mode only
        exam.keys = {c: k for c, k in exam.keys.items() if c == ALL_CODES}
    else:
        exam.keys = {c: k for c, k in exam.keys.items() if c != ALL_CODES}
    exam.keys.update(keys)
    if isinstance(exam.in_order, list):
        exam.in_order = {c: list(exam.in_order) for c in exam.keys if c not in keys}
    else:
        exam.in_order = {c: v for c, v in exam.in_order.items() if c in exam.keys and c not in keys}
    exam.in_order.update(in_order)
    exam.code_column = code_column
    weights = data.get("weights") or {}
    exam.weights = {part_of(p): str(w) for p, w in weights.items() if part_of(p) in sections and str(w).strip()}
    scoring = data.get("scoring") or {}
    try:
        for f in ("correct", "incorrect", "unmarked", "kkm"):
            if f in scoring:
                setattr(exam, f, float(scoring[f]))
    except (TypeError, ValueError):
        warnings.append("Aturan skor di file ini tidak terbaca, skor tidak diubah.")
    if "scale_100" in scoring:
        exam.scale_100 = bool(scoring["scale_100"])
    return warnings


# ---------------------------------------------------------------------------
# Results and scoring
# ---------------------------------------------------------------------------


def current(out, pattern):
    """out.rglob(pattern) without anything under an archive folder."""
    return [f for f in out.rglob(pattern) if ARCHIVE_DIR not in f.relative_to(out).parts]


def read_results(output_dir):
    """Rows from every Results_*.csv under output_dir (newest wins per file),
    plus rows of sheets that could not be read (Manual/ErrorFiles.csv).
    A sheet from a subfolder of the scans gets that folder in its file_id
    ("7B/scan_0001.jpg"): scanners restart their numbering in every folder,
    and corrections are keyed by file_id."""
    out = Path(output_dir)
    rows, errors = {}, {}
    if not out.is_dir():
        return [], []

    def sheet(r, f):
        sub = f.parent.parent.relative_to(out).as_posix()
        return r if sub == "." else {**r, "file_id": f"{sub}/{r['file_id']}"}

    files = sorted(current(out, "Results/Results_*.csv"), key=lambda p: p.stat().st_mtime)
    for f in files:
        with open(f, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                rows[(f.parent.parent, r["file_id"])] = sheet(r, f)
    for f in current(out, "Manual/ErrorFiles.csv"):
        with open(f, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                errors[(f.parent.parent, r["file_id"])] = sheet(r, f)
    for k in rows:
        errors.pop(k, None)
    return list(rows.values()), list(errors.values())


def done_sheets(output_dir):
    """[folder relative to the scans, file_id] of every sheet already read
    (or found unreadable) in output_dir: what a resumed run can skip."""
    out = Path(output_dir)
    if not out.is_dir():
        return []
    done = set()
    for pattern in ("Results/Results_*.csv", "Manual/ErrorFiles.csv"):
        for f in current(out, pattern):
            rel = f.parent.parent.relative_to(out).as_posix()
            with open(f, newline="", encoding="utf-8") as fh:
                done.update((rel, r["file_id"]) for r in csv.DictReader(fh) if r.get("file_id"))
    return sorted([list(k) for k in done])


def count_sheets(paths):
    """Number of sheets in the scan files: a PDF counts once per page."""
    total = 0
    for p in paths:
        if Path(p).suffix.lower() == ".pdf":
            try:
                import fitz

                with fitz.open(str(p)) as doc:
                    total += max(1, doc.page_count)
                continue
            except Exception:  # unreadable PDF: the run will report it
                pass
        total += 1
    return total


@dataclass
class SheetResult:
    file_id: str
    image: str
    answers: dict  # question label -> value read (after corrections)
    info: dict  # identity columns read from the sheet (after corrections)
    code: str = ""
    correct: int = 0
    wrong: int = 0
    blank: int = 0
    multi: int = 0
    ignored: int = 0  # answers not counted (Exam.in_order: after a blank)
    score: float = None
    grade: float = None
    passed: bool = None
    rank: int = None  # 1 = highest grade (rank_results)
    parts: dict = field(default_factory=dict)  # part -> {benar, salah, kosong, diabaikan, skor, nilai}
    flags: list = field(default_factory=list)
    skipped: list = field(default_factory=list)  # question labels of `ignored`
    student: dict = None  # matched row of the student list
    unreadable: bool = False
    corrected: bool = False
    raw: dict = field(default_factory=dict)  # values as read from the scan

    @property
    def has_nis(self):
        return "NIS" in self.info or "NoPeserta" in self.info

    @property
    def nis(self):
        return self.info.get("NIS") or self.info.get("NoPeserta") or ""

    @property
    def name(self):
        if self.student and self.student.get("nama"):
            return self.student["nama"]
        return self.info.get("Nama", "").strip()

    @property
    def kelas(self):
        return (self.student or {}).get("kelas", "")


def normalize_name(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().upper()


def normalize_nis(value):
    v = re.sub(r"\s+", "", str(value or ""))
    return v.lstrip("0") or v


def as_sections(questions):
    return questions if isinstance(questions, dict) else {"": list(questions)}


def keys_for(exam, code):
    """{part: key string} for an exam code, or None when it has no key."""
    k = exam.keys.get(code)
    if k is None:
        return None
    return {"": k} if isinstance(k, str) else k


def active_codes(exam):
    """Codes whose keys are scored: the per-code keys when the exam reads a
    code column, otherwise ALL_CODES. The keys of the other mode stay in
    exam.keys, so switching the mode back and forth loses nothing."""
    if exam.code_column:
        return [c for c in exam.keys if c != ALL_CODES]
    return [ALL_CODES] if ALL_CODES in exam.keys else []


def in_order_for(exam, code):
    """Parts of `code` that must be answered in order."""
    if isinstance(exam.in_order, list):  # v1.3.x: one setting for every code
        return exam.in_order
    return exam.in_order.get(code, [])


def set_in_order(exam, code, part, on):
    """Turn 'must be answered in order' on or off for one part of one code."""
    if isinstance(exam.in_order, list):  # from now on per code
        exam.in_order = {c: list(exam.in_order) for c in {*exam.keys, ALL_CODES}}
    parts = [p for p in exam.in_order.get(code, []) if p != part]
    exam.in_order[code] = parts + [part] if on else parts


def uses_in_order(exam):
    """Whether any scored code has a part that must be answered in order
    (the rekap then gets a 'Tidak dihitung' column)."""
    return any(in_order_for(exam, c) for c in active_codes(exam))


def question_label(section, q):
    """Readable label: 'I-5' for part I question 5, 'q5' -> '5'."""
    num = q[len(q.rstrip("0123456789")):]
    return f"{section}-{num}" if section else num


def score_results(rows, errors, exam, questions, students=None, blank_source=""):
    """`questions`: {part: labels} (or a plain list for a single part).
    `blank_source`: file the template was made from; hidden when it was read
    back with nothing marked (the blank sheet lying among the scans)."""
    sections = as_sections(questions)
    all_q = [q for labels in sections.values() for q in labels]
    if blank_source:
        stem = Path(blank_source).stem
        skip_keys = {"file_id", "input_path", "output_path", "score"}
        rows = [r for r in rows if Path(r["file_id"]).stem != stem
                or any(v.strip() for k, v in r.items() if k not in skip_keys)]  # fmt: skip
    _upgrade_corrections(exam, rows)
    students = students or []
    by_nis, by_name = {}, {}
    for s in students:
        if s.get("nis"):
            by_nis.setdefault(normalize_nis(s["nis"]), s)
        if s.get("nama"):
            by_name.setdefault(normalize_name(s["nama"]), s)

    results = []
    for r in rows:
        fid = r["file_id"]
        fixes = exam.corrections.get(fid, {})
        values = {**r, **fixes}
        answers = {q: values.get(q, "") for q in all_q}
        skip = {"file_id", "input_path", "output_path", "score"}
        info = {k: v for k, v in values.items() if k not in skip and k not in answers}
        res = SheetResult(fid, r.get("output_path") or r.get("input_path", ""), answers, info)
        res.corrected = bool(fixes)
        res.raw = dict(r)
        res.code = values.get(exam.code_column, "") if exam.code_column else ALL_CODES
        _score_one(res, exam, sections)
        if students:
            if res.has_nis:
                res.student = by_nis.get(normalize_nis(res.nis))
            else:  # sheet without a number field: match on the name
                res.student = by_name.get(normalize_name(res.info.get("Nama", "")))
        results.append(res)

    for e in errors:
        res = SheetResult(e["file_id"], e.get("input_path", ""), {q: "" for q in all_q}, {})
        res.unreadable = True
        res.flags.append("LJK tidak terbaca (kotak sudut/halaman tidak ditemukan)")
        results.append(res)

    # identity checks across sheets
    seen = {}
    for res in results:
        if res.unreadable:
            continue
        if not res.has_nis:
            if students and res.student is None:
                res.flags.append("Nama tidak ada di daftar peserta")
            continue
        nis = normalize_nis(res.nis)
        if not res.nis.strip():
            res.flags.append("No. peserta kosong")
        else:
            seen.setdefault(nis, []).append(res)
        if students and res.nis.strip() and res.student is None:
            res.flags.append(NOT_LISTED + _guess_participant(res, students, results))
    for dup in (v for v in seen.values() if len(v) > 1):
        for res in dup:
            res.flags.append("No. peserta dobel: " + ", ".join(o.file_id for o in dup if o is not res))
    results.sort(key=lambda x: (not x.flags, natural_key(x.kelas), natural_key(x.name), x.file_id))
    return results


NOT_LISTED = "No. peserta tidak ada di daftar peserta"


def _close_numbers(a, b):
    """Whether two numbers differ by one digit (wrong, missing or extra) or
    by two neighbouring digits swapped: a likely mis-marked number."""
    if a == b:
        return False
    if len(a) == len(b):
        diff = [i for i in range(len(a)) if a[i] != b[i]]
        return len(diff) == 1 or (len(diff) == 2 and diff[1] == diff[0] + 1
                                  and a[diff[0]] == b[diff[1]] and a[diff[1]] == b[diff[0]])  # fmt: skip
    short, long_ = sorted((a, b), key=len)
    return len(long_) - len(short) == 1 and any(long_[:i] + long_[i + 1:] == short for i in range(len(long_)))


def _guess_participant(res, students, results):
    """' (mungkin <No. Peserta> <Nama>?)' for a sheet whose number is not in
    the list: the participant without a sheet whose name matches the name on
    the sheet, or else the only one whose number is one slip away."""
    taken = {id(r.student) for r in results if r.student is not None}
    free = [s for s in students if id(s) not in taken and s.get("nis")]
    name = normalize_name(res.info.get("Nama", ""))
    found = [s for s in free if name and normalize_name(s.get("nama")) == name]
    if not found:
        nis = normalize_nis(res.nis)
        found = [s for s in free if _close_numbers(nis, normalize_nis(s["nis"]))]
    if len(found) != 1:
        return ""
    s = found[0]
    return f" (mungkin {s['nis']}{' ' + s['nama'] if s.get('nama') else ''}?)"


def duplicate_numbers(students):
    """No. Peserta that appear more than once in the participant list."""
    seen, dup = set(), []
    for s in students:
        n = normalize_nis(s.get("nis"))
        if n and n in seen and s["nis"] not in dup:
            dup.append(s["nis"])
        seen.add(n)
    return dup


def _upgrade_corrections(exam, rows):
    """Corrections saved by v1.3.0 and older are keyed by the bare file name,
    also for sheets in a subfolder: move them to the sheet's full file_id when
    that name belongs to one sheet only."""
    names = {}
    for r in rows:
        bare = r["file_id"].rsplit("/", 1)[-1]
        names[bare] = names.get(bare, 0) + 1
    for r in rows:
        fid = r["file_id"]
        bare = fid.rsplit("/", 1)[-1]
        if fid != bare and fid not in exam.corrections and bare in exam.corrections and names[bare] == 1:
            exam.corrections[fid] = exam.corrections.pop(bare)


def _score_one(res, exam, sections):
    multi = [question_label(sec, q) for sec, labels in sections.items() for q in labels if len(res.answers[q]) > 1]
    if multi:
        res.flags.append("Jawaban ganda: " + ", ".join(multi))
    res.multi = len(multi)
    keys = keys_for(exam, res.code)
    if keys is None:
        if active_codes(exam):
            res.flags.append(f"Kode naskah '{res.code}' tidak punya kunci jawaban")
        return
    total_best = 0.0
    for sec, labels in sections.items():
        key = keys.get(sec, "")
        if not key:
            continue
        weights, _ = parse_weights(exam.weights.get(sec, ""))
        part = {"benar": 0, "salah": 0, "kosong": 0, "diabaikan": 0}
        score = best = 0.0
        gap = None  # first blank question when the part must be answered in order
        for n, (q, k) in enumerate(zip(labels, key), 1):
            a = res.answers[q]
            points = weights.get(n, exam.correct)
            best += points
            if a and gap is not None:
                part["diabaikan"] += 1
                res.skipped.append(q)
                score += exam.unmarked
            elif a == "":
                if sec in in_order_for(exam, res.code) and gap is None:
                    gap = question_label(sec, q)
                part["kosong"] += 1
                score += exam.unmarked
            elif a == k:
                part["benar"] += 1
                score += points
            else:
                part["salah"] += 1
                score += exam.incorrect
        if part["diabaikan"]:
            res.flags.append(f"Soal {gap} kosong: {part['diabaikan']} jawaban sesudahnya tidak dihitung (harus berurutan)")
        part["skor"] = round(score, 2)
        part["nilai"] = round(part["skor"] / best * 100, 2) if exam.scale_100 and best else part["skor"]
        res.parts[sec] = part
        total_best += best
        res.correct += part["benar"]
        res.wrong += part["salah"]
        res.blank += part["kosong"]
        res.ignored += part["diabaikan"]
    if not res.parts:
        return
    res.score = round(sum(p["skor"] for p in res.parts.values()), 2)
    res.grade = round(res.score / total_best * 100, 2) if exam.scale_100 and total_best else res.score
    if exam.kkm:
        res.passed = res.grade >= exam.kkm


RANK_ORDERS = {
    "nilai": "Nilai tertinggi dulu",
    "nama": "Kelompok, lalu nama",
    "file": "Nama file scan",
}


def rank_results(results):
    """Rank by grade, highest = 1; equal grades share a rank (1, 2, 2, 4).
    Sheets without a grade get None."""
    graded = sorted((r for r in results if r.grade is not None), key=lambda r: -r.grade)
    prev, rank = None, 0
    for i, r in enumerate(graded, 1):
        if r.grade != prev:
            rank, prev = i, r.grade
        r.rank = rank
    for r in results:
        if r.grade is None:
            r.rank = None
    return results


def order_results(results, order="nama"):
    """The rekap's row order (RANK_ORDERS); ungraded or unreadable sheets go last."""
    rank_results(results)
    if order == "nilai":
        key = lambda r: (r.rank is None, r.rank or 0, natural_key(r.name), r.file_id)  # noqa: E731
    elif order == "file":
        key = lambda r: natural_key(r.file_id)  # noqa: E731
    else:
        key = lambda r: (r.unreadable, natural_key(r.kelas), natural_key(r.name), natural_key(r.file_id))  # noqa: E731
    return sorted(results, key=key)


def absent_students(results, students):
    present = {id(r.student) for r in results if r.student is not None}
    return [s for s in students if id(s) not in present]


def summary(results):
    grades = [r.grade for r in results if r.grade is not None]
    passed = [r for r in results if r.passed]
    return {
        "peserta": len(results),
        "dinilai": len(grades),
        "rata": round(sum(grades) / len(grades), 2) if grades else None,
        "tertinggi": max(grades) if grades else None,
        "terendah": min(grades) if grades else None,
        "tuntas": len(passed),
        "perlu_dicek": sum(1 for r in results if r.flags),
    }


# ---------------------------------------------------------------------------
# Student list
# ---------------------------------------------------------------------------

COLUMN_HINTS = {
    "nis": ("no. peserta", "no peserta", "nomor peserta", "nopeserta", "nis", "nisn", "nip", "nik",
            "no induk", "nomor induk", "id peserta", "id karyawan"),
    "nama": ("nama", "name"),
    "kelas": ("kelompok", "kelas", "rombel", "class", "divisi", "bagian", "unit", "departemen", "ruang"),
}


def guess_student_columns(headers):
    found = {}
    for key, hints in COLUMN_HINTS.items():
        for h in headers:
            low = str(h).strip().lower()
            if any(low == hint or low.startswith(hint) for hint in hints):
                found[key] = h
                break
    return found


def read_table(path):
    """(headers, rows as dicts of strings) from .xlsx/.xls/.csv."""
    import pandas as pd

    path = Path(path)
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, dtype=str, sep=None, engine="python")
    else:
        df = pd.read_excel(path, dtype=str)
    df = df.fillna("")
    df.columns = [str(c).strip() for c in df.columns]
    return list(df.columns), df.to_dict("records")


def load_students(path, columns=None):
    headers, rows = read_table(path)
    columns = columns or guess_student_columns(headers)
    if not columns.get("nis") and not columns.get("nama"):
        raise ValueError("Kolom No. Peserta atau Nama tidak ditemukan di daftar peserta")
    students = []
    for r in rows:
        s = {k: str(r.get(h, "")).strip() for k, h in columns.items() if h}
        if s.get("nis") or s.get("nama"):
            students.append(s)
    return students, headers, columns


# ---------------------------------------------------------------------------
# Item analysis
# ---------------------------------------------------------------------------


MIN_DISCRIMINATION = 5  # sheets needed for a discrimination figure
MIN_ITEM_NOTES = 10  # sheets needed before a question is pointed at


def _item_groups(results, exam):
    """[(code, graded results)]: one group per exam code when the exam reads a
    code column and has more than one key (the codes may order the questions
    differently), otherwise one group ("") with every graded sheet."""
    graded = [r for r in results if r.grade is not None]
    codes = active_codes(exam)
    if not exam.code_column or len(codes) < 2:
        return [("", graded)]
    groups = [(c, [r for r in graded if r.code == c]) for c in sorted(codes, key=natural_key)]
    return [(c, rs) for c, rs in groups if rs]


def _top_size(group):
    return max(1, round(len(group) * 0.27))


def discrimination(group, correct):
    """Daya beda: share correct in the top 27% (by grade) minus the bottom
    27%, from -1 to 1. None with too few sheets to say."""
    if len(group) < MIN_DISCRIMINATION:
        return None
    ranked = sorted(group, key=lambda r: -r.grade)
    n = _top_size(ranked)
    top = sum(correct(r) for r in ranked[:n]) / n
    bottom = sum(correct(r) for r in ranked[-n:]) / n
    return round(top - bottom, 2)


def item_analysis(results, exam, questions):
    """One row per question (per exam code when the codes have their own
    keys): % correct, difficulty, discrimination, the answer chosen most and
    how often each option was chosen. `catatan` points at questions worth a
    second look: a key that may be mistyped, or a negative discrimination."""
    sections = as_sections(questions)
    rows = []
    for code, group in _item_groups(results, exam):
        for sec, labels in sections.items():
            for i, q in enumerate(labels):
                counts = {o: 0 for o in OPTIONS_ORDER}
                counts.update({"Kosong": 0, "Ganda": 0})
                answered = {}  # result -> answer counted (not counted = blank) is right
                given = {}  # result -> that answer
                for r in group:
                    key = (keys_for(exam, r.code) or {}).get(sec, "")
                    if i >= len(key):
                        continue
                    a = "" if q in r.skipped else r.answers[q]
                    if a == "":
                        counts["Kosong"] += 1
                    elif len(a) > 1:
                        counts["Ganda"] += 1
                    else:
                        counts[a] = counts.get(a, 0) + 1
                    answered[id(r)] = a == key[i]
                    given[id(r)] = a
                if not answered:
                    continue
                pct = round(sum(answered.values()) / len(answered) * 100, 1)
                level = "Mudah" if pct > 70 else "Sedang" if pct >= 30 else "Sukar"
                per_code = {}
                for c in [code] if code else active_codes(exam):
                    k = (keys_for(exam, c) or {}).get(sec, "")
                    if i < len(k):
                        per_code[c] = k[i]
                key_text = per_code.get(code or ALL_CODES) or ", ".join(f"{c}: {v}" for c, v in sorted(per_code.items()))
                chosen = {o: n for o, n in counts.items() if o not in ("Kosong", "Ganda") and n}
                top = max(chosen.values(), default=0)
                most = "/".join(o for o, n in chosen.items() if n == top)
                sheets = [r for r in group if id(r) in answered]
                power = discrimination(sheets, lambda r: answered[id(r)])
                notes = []
                if len(sheets) >= MIN_ITEM_NOTES and len(per_code) == 1 and len(most) == 1 and most != key_text:
                    # a mistyped key: the best sheets mostly agree on another answer
                    best = sorted(sheets, key=lambda r: -r.grade)[: _top_size(sheets)]
                    picks = [given[id(r)] for r in best]
                    if picks.count(most) * 2 >= len(picks) and picks.count(most) > picks.count(key_text):
                        notes.append(f"Cek kunci: kebanyakan peserta, juga yang nilainya tertinggi, memilih {most}, "
                                     f"bukan {key_text}")  # fmt: skip
                if len(sheets) >= MIN_ITEM_NOTES and power is not None and power <= -0.2:
                    notes.append("Daya beda negatif: peserta bernilai rendah lebih banyak yang benar")
                weights, _ = parse_weights(exam.weights.get(sec, ""))
                rows.append({"kode": code, "bagian": sec, "soal": question_label(sec, q), "kunci": key_text,
                             "bobot": weights.get(i + 1, exam.correct), "benar": pct, "kategori": level,
                             "daya_beda": power, "terbanyak": most, "catatan": "; ".join(notes),
                             **counts})  # fmt: skip
    return rows


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


def rekap_table(results, exam, named):
    """Headers and rows of the "Rekap Nilai" sheet (also the CSV export),
    `results` already in rekap order; `named`: the template's named parts."""
    id_cols = [c for c in _identity_columns(results) if c not in ("NIS", "NoPeserta", "Nama", exam.code_column)]
    headers = ["No", "No. Peserta", "Nama", "Kelompok", *id_cols]
    if exam.code_column:
        headers.append("Kode Naskah")
    for sec in named:
        headers += [f"Benar {sec}", f"Nilai {sec}"]
    headers += ["Benar", "Salah", "Kosong"] + (["Tidak dihitung"] if uses_in_order(exam) else [])
    headers += ["Skor", "Nilai", "Peringkat"]
    if exam.kkm:
        headers.append("Keterangan")
    headers += ["Catatan", "File"]
    rows = []
    for i, r in enumerate(results, 1):
        row = [i, r.nis, r.name, r.kelas, *[r.info.get(c, "") for c in id_cols]]
        if exam.code_column:
            row.append(r.code)
        for sec in named:
            part = r.parts.get(sec, {})
            row += [part.get("benar"), part.get("nilai")]
        row += [r.correct, r.wrong, r.blank] + ([r.ignored] if uses_in_order(exam) else [])
        row += [r.score, r.grade, r.rank]
        if exam.kkm:
            row.append("" if r.passed is None else "Lulus" if r.passed else "Tidak lulus")
        row += ["; ".join(r.flags), r.file_id]
        rows.append(row)
    return headers, rows


def export_table(results, exam, questions):
    """Every exportable column: the "Rekap Nilai" columns, then the answer
    to every question, rows in rekap order."""
    sections = as_sections(questions)
    results = order_results(results, exam.rekap_order)
    headers, rows = rekap_table(results, exam, [s for s in sections if s])
    labels = [(sec, q) for sec, qs in sections.items() for q in qs]
    headers = headers + [f"Jawaban {question_label(sec, q)}" for sec, q in labels]
    rows = [row + [r.answers[q] for _, q in labels] for r, row in zip(results, rows)]
    return headers, rows


def _write_csv(path, headers, rows):
    with Path(path).open("w", newline="", encoding="utf-8-sig") as f:  # BOM: Excel reads the accents right
        w = csv.writer(f)
        w.writerow(headers)
        for row in rows:
            w.writerow(["" if v is None else v for v in row])


def export_csv(path, results, exam, questions):
    """The rekap as one CSV table (UTF-8, comma separated) for other
    programs: the "Rekap Nilai" columns, then the answer to every question."""
    path = Path(path)
    _write_csv(path, *export_table(results, exam, questions))
    return path


def export_custom(path, results, exam, questions, columns):
    """The columns the user picked, in their order and under their own
    headers: `columns` is [(column, header), ...]. One plain table with the
    headers on row 1 (.csv, otherwise Excel). Columns this exam does not
    have are left out; returns (path, the left-out columns)."""
    headers, rows = export_table(results, exam, questions)
    index = {h: i for i, h in enumerate(headers)}
    picked = [(index[c], t or c) for c, t in columns if c in index]
    missing = [c for c, _ in columns if c not in index]
    out_headers = [t for _, t in picked]
    out_rows = [[row[i] for i, _ in picked] for row in rows]
    path = Path(path)
    if path.suffix.lower() == ".csv":
        _write_csv(path, out_headers, out_rows)
        return path, missing
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    ws.append(out_headers)
    for row in out_rows:
        ws.append(row)
    for c, h in enumerate(out_headers, 1):
        ws.cell(1, c).font = Font(bold=True)
        longest = max([len(str(h))] + [len(str(v[c - 1])) for v in out_rows if v[c - 1] is not None])
        ws.column_dimensions[get_column_letter(c)].width = min(45, max(6, longest + 2))
    ws.freeze_panes = "A2"
    wb.save(path)
    return path, missing


EXPORT_FORMATS_FILE = "format_ekspor.json"  # saved "Ekspor lanjutan" formats, next to the session file
LAST_EXPORT_FORMAT = ""  # the settings used last time, kept under this name


def export_formats_path(session=None):
    return Path(session or Exam.session_path()).with_name(EXPORT_FORMATS_FILE)


def load_export_formats(session=None):
    """{name: {"kolom": [[column, header], ...], "jenis": "xlsx" | "csv"}}"""
    try:
        data = json.loads(export_formats_path(session).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    formats = {}
    for name, fmt in (data if isinstance(data, dict) else {}).items():
        if isinstance(fmt, dict) and isinstance(fmt.get("kolom"), list):
            cols = [[str(c[0]), str(c[1])] for c in fmt["kolom"] if isinstance(c, list) and len(c) == 2]
            formats[str(name)] = {"kolom": cols, "jenis": "csv" if fmt.get("jenis") == "csv" else "xlsx"}
    return formats


def save_export_formats(formats, session=None):
    path = export_formats_path(session)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(formats, indent=1, ensure_ascii=False), encoding="utf-8")


def export_excel(path, results, exam, questions, students=None, title=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sections = as_sections(questions)
    named = [s for s in sections if s]
    results = order_results(results, exam.rekap_order)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    red = PatternFill("solid", fgColor="F8CBAD")
    grey = PatternFill("solid", fgColor="EDEDED")
    yellow = PatternFill("solid", fgColor="FFE699")
    wb = Workbook()

    def table(ws, start_row, headers, rows, widths=None):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(start_row, c, h)
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r, values in enumerate(rows, start_row + 1):
            for c, v in enumerate(values, 1):
                ws.cell(r, c, v)
        for c, h in enumerate(headers, 1):
            longest = max([len(str(h))] + [len(str(v[c - 1])) for v in rows if c - 1 < len(v)])
            ws.column_dimensions[get_column_letter(c)].width = (
                widths.get(h) if widths and h in widths else min(45, max(6, longest + 2))
            )
        ws.freeze_panes = ws.cell(start_row + 1, 1)

    # Rekap nilai
    ws = wb.active
    ws.title = "Rekap Nilai"
    stats = summary(results)
    ws["A1"] = title or exam.name or "Rekap Nilai Ujian"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Peserta: {stats['peserta']}   Rata-rata: {stats['rata']}   "
        f"Tertinggi: {stats['tertinggi']}   Terendah: {stats['terendah']}"
        + (f"   Lulus (batas {exam.kkm:g}): {stats['tuntas']}" if exam.kkm else "")
    )
    ws["A3"] = f"Diekspor: {time.strftime('%d-%m-%Y %H:%M')}"
    headers, rows = rekap_table(results, exam, named)
    table(ws, 5, headers, rows, {"Catatan": 50})
    for i, r in enumerate(results):
        if r.flags:
            for c in range(1, len(headers) + 1):
                ws.cell(6 + i, c).fill = red

    # Jawaban per soal
    ws = wb.create_sheet("Jawaban")
    labels = [(sec, q) for sec, qs in sections.items() for q in qs]
    lead = ["No. Peserta", "Nama"] + (["Kode Naskah"] if exam.code_column else [])
    headers = lead + [question_label(sec, q) for sec, q in labels]
    rows = [[r.nis, r.name] + ([r.code] if exam.code_column else []) + [r.answers[q] for _, q in labels]
            for r in results]  # fmt: skip
    table(ws, 1, headers, rows)
    if uses_in_order(exam):
        ws.freeze_panes = None
        ws.insert_rows(1)
        ws["A1"] = "Kuning: jawaban tidak dihitung karena ada soal kosong sebelumnya (harus diisi berurutan)"
        ws.freeze_panes = "A3"
    top = 3 if uses_in_order(exam) else 2
    for i, r in enumerate(results):
        keys = keys_for(exam, r.code) or {}
        for j, (sec, q) in enumerate(labels):
            n = sections[sec].index(q)
            key = keys.get(sec, "")
            cell = ws.cell(top + i, len(lead) + 1 + j)
            if r.answers[q] == "":
                cell.fill = grey
            elif q in r.skipped:
                cell.fill = yellow
            elif n < len(key) and r.answers[q] != key[n]:
                cell.fill = red

    # Analisis butir
    ws = wb.create_sheet("Analisis Butir")
    analysis = item_analysis(results, exam, sections)
    by_code = any(a["kode"] for a in analysis)
    counted = OPTIONS_ORDER + ["Kosong", "Ganda"]
    headers = (["Kode Naskah"] if by_code else []) + ["Soal", "Kunci", "Bobot", "% Benar", "Kategori",
                                                      "Daya Beda", "Terbanyak"] + counted + ["Catatan"]  # fmt: skip
    table(ws, 1, headers, [([a["kode"]] if by_code else []) + [a["soal"], a["kunci"], a["bobot"], a["benar"],
                           a["kategori"], a["daya_beda"], a["terbanyak"]] + [a[o] for o in counted] + [a["catatan"]]
                           for a in analysis], {"Catatan": 60})  # fmt: skip
    for i, a in enumerate(analysis, 2):
        if a["catatan"]:
            for c in range(1, len(headers) + 1):
                ws.cell(i, c).fill = yellow
    end = len(analysis) + 3
    ws.cell(end, 1, "Daya Beda: selisih persentase benar kelompok nilai tertinggi (27%) dan terendah (27%), "
                    "dari -1 sampai 1. Di atas 0,3 baik; di bawah 0,2 soal kurang membedakan; "
                    f"negatif perlu dicek. Dihitung bila ada minimal {MIN_DISCRIMINATION} LJK.")
    ws.cell(end + 1, 1, f"Kuning: soal yang perlu dicek (kunci mungkin salah ketik, atau daya beda negatif); "
                        f"ditandai bila ada minimal {MIN_ITEM_NOTES} LJK.")

    if students:
        ws = wb.create_sheet("Tidak Ikut Ujian")
        missing = absent_students(results, students)
        table(ws, 1, ["No", "No. Peserta", "Nama", "Kelompok"],
              [[i, s.get("nis", ""), s.get("nama", ""), s.get("kelas", "")]
               for i, s in enumerate(missing, 1)])  # fmt: skip

    path = Path(path)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Per-participant report (PDF)
# ---------------------------------------------------------------------------

REPORT_CSS = """
* {font-family: sans-serif; font-size: 10pt;}
h1 {font-size: 15pt; margin: 0 0 2pt 0;}
h2 {font-size: 11pt; margin: 10pt 0 3pt 0;}
p.muted {color: #666666; margin: 0 0 8pt 0;}
table {border-collapse: collapse; width: 100%;}
td, th {border: 1px solid #b0b7c3; padding: 2pt 4pt; text-align: center;}
th {background-color: #ddebf7;}
td.l {text-align: left;}
td.ok {background-color: #d9f2df;}
td.bad {background-color: #f8cbad;}
td.blank {background-color: #ededed;}
td.skip {background-color: #ffe699; text-decoration: line-through;}
"""


def _report_html(r, exam, sections, title, total):
    from html import escape as esc

    keys = keys_for(exam, r.code) or {}
    ident = [("No. Peserta", r.nis), ("Nama", r.name), ("Kelompok", r.kelas)]
    ident += [("Kode naskah" if k == exam.code_column else k, v)
              for k, v in r.info.items() if k not in ("NIS", "NoPeserta", "Nama")]  # fmt: skip
    ident.append(("File scan", r.file_id))
    rows = "".join(f'<tr><td class="l" width="30%">{esc(k)}</td><td class="l">{esc(str(v))}</td></tr>'
                   for k, v in ident if str(v).strip())  # fmt: skip
    status = "" if r.passed is None else ("Lulus" if r.passed else "Tidak lulus")
    rank = f"{r.rank} dari {total}" if r.rank else "-"
    score = (
        "<tr><th>Benar</th><th>Salah</th><th>Kosong</th>"
        + ("<th>Tidak dihitung</th>" if uses_in_order(exam) else "") + "<th>Skor</th><th>Nilai</th><th>Peringkat</th>"
        + ("<th>Keterangan</th>" if exam.kkm else "") + "</tr>"
        + f"<tr><td>{r.correct}</td><td>{r.wrong}</td><td>{r.blank}</td>"
        + (f"<td>{r.ignored}</td>" if uses_in_order(exam) else "")
        + f"<td>{'' if r.score is None else f'{r.score:g}'}</td><td><b>{'' if r.grade is None else f'{r.grade:g}'}</b></td>"
        + f"<td>{rank}</td>" + (f"<td>{status}</td>" if exam.kkm else "") + "</tr>"
    )
    parts = ""
    for sec, labels in sections.items():
        key = keys.get(sec, "")
        cells = []
        for n, q in enumerate(labels):
            a = r.answers.get(q, "")
            k = key[n] if n < len(key) else ""
            css = "blank" if a == "" else "skip" if q in r.skipped else ("ok" if a == k else "bad") if k else ""
            cells.append(f'<td>{esc(str(n + 1))}</td><td class="{css}">{esc(a or "-")}</td><td>{esc(k or "")}</td>')
        per_row = 4
        body = "".join("<tr>" + "".join(cells[i : i + per_row]) + "</tr>" for i in range(0, len(cells), per_row))
        head = "<tr>" + "<th>No</th><th>Jawab</th><th>Kunci</th>" * min(per_row, len(cells)) + "</tr>"
        part = r.parts.get(sec)
        summary = f" (benar {part['benar']}, nilai {part['nilai']:g})" if part else ""
        parts += f"<h2>{esc(f'Bagian {sec}' if sec else 'Jawaban')}{esc(summary)}</h2><table>{head}{body}</table>"
    notes = f"<h2>Catatan</h2><p>{esc('; '.join(r.flags))}</p>" if r.flags else ""
    return (
        f"<h1>{esc(title or exam.name or 'Hasil Ujian')}</h1>"
        f'<p class="muted">Laporan hasil peserta • dicetak {time.strftime("%d-%m-%Y %H:%M")}</p>'
        f"<table>{rows}</table><h2>Nilai</h2><table>{score}</table>{parts}{notes}"
    )


def export_reports_pdf(path, results, exam, questions, title="", separate=False):
    """One A4 page (or more) per graded sheet: identity, score, rank and every
    answer next to the key. Unreadable sheets are left out. `separate`: `path`
    is a folder that gets one PDF per participant ("<No. Peserta> <Nama>.pdf",
    to hand out one by one); returns the files written."""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    sections = as_sections(questions)
    results = order_results([r for r in results if not r.unreadable], exam.rekap_order)
    total = sum(1 for r in results if r.rank)
    page = fitz.paper_rect("a4")
    where = page + (36, 36, -36, -36)

    def write(target, sheets):
        writer = fitz.DocumentWriter(str(target))
        try:
            for r in sheets:
                story = fitz.Story(html=_report_html(r, exam, sections, title, total), user_css=REPORT_CSS)
                more = True
                while more:
                    dev = writer.begin_page(page)
                    more, _ = story.place(where)
                    story.draw(dev)
                    writer.end_page()
        finally:
            writer.close()

    path = Path(path)
    if not separate:
        write(path, results)
        return path
    path.mkdir(parents=True, exist_ok=True)
    files, used = [], set()
    for r in results:
        stem = safe_name(" ".join(p for p in (r.nis.strip(), r.name) if p) or Path(r.file_id).stem)
        name, n = stem, 1
        while name.lower() in used:  # same number and name twice: keep both
            n += 1
            name = f"{stem} ({n})"
        used.add(name.lower())
        write(path / f"{name}.pdf", [r])
        files.append(path / f"{name}.pdf")
    return files


def _identity_columns(results):
    cols = []
    for r in results:
        for c in r.info:
            if c not in cols:
                cols.append(c)
    return cols
