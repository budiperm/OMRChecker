"""Tests for ujian.py: scoring with per-exam-code keys, corrections, student
list, item analysis, Excel export, template library, and the wizard's
processing path (template from the library, scan folder untouched) on the
real LJK.

Run: python -m pytest tests_gui
"""

import json
import random
import shutil
import subprocess
import sys
from pathlib import Path

import cv2
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(Path(__file__).parent)]

import synth_ljk as synth  # noqa: E402
import template_builder as tb  # noqa: E402
import ujian  # noqa: E402
from testdata import needs  # noqa: E402

Q5 = [f"q{i}" for i in range(1, 6)]


def row(fid, answers, **info):
    r = {"file_id": fid, "input_path": f"/scan/{fid}", "output_path": f"/out/{fid}", "score": "0"}
    r.update({q: a for q, a in zip(Q5, answers)})
    r.update(info)
    return r


def exam(**kw):
    e = ujian.Exam(code_column="KodeSoal", keys={"101": "ABCDE", "102": "EDCBA"})
    for k, v in kw.items():
        setattr(e, k, v)
    return e


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------


def test_each_code_uses_its_own_key():
    rows = [
        row("a.jpg", ["A", "B", "C", "D", "E"], KodeSoal="101", NIS="1", Nama="ANI"),
        row("b.jpg", ["A", "B", "C", "D", "E"], KodeSoal="102", NIS="2", Nama="BUDI"),
    ]
    res = {r.file_id: r for r in ujian.score_results(rows, [], exam(), Q5)}
    assert (res["a.jpg"].correct, res["a.jpg"].grade) == (5, 100)
    assert (res["b.jpg"].correct, res["b.jpg"].wrong, res["b.jpg"].grade) == (1, 4, 20)


def test_blank_multi_unknown_code_and_negative_marking():
    rows = [
        row("a.jpg", ["A", "", "AC", "D", "B"], KodeSoal="101", NIS="1"),
        row("c.jpg", ["A", "B", "C", "D", "E"], KodeSoal="999", NIS="3"),
    ]
    e = exam(correct=4, incorrect=-1, unmarked=0, kkm=50)
    res = {r.file_id: r for r in ujian.score_results(rows, [], e, Q5)}
    a = res["a.jpg"]
    assert (a.correct, a.wrong, a.blank, a.multi) == (2, 2, 1, 1)
    assert a.score == 2 * 4 - 2 and a.grade == 30 and a.passed is False
    assert any("ganda" in f for f in a.flags)
    c = res["c.jpg"]
    assert c.grade is None and any("999" in f for f in c.flags)


def test_single_key_for_everyone():
    e = ujian.Exam(keys={ujian.ALL_CODES: "ABC"})  # key shorter than the sheet
    res = ujian.score_results([row("a.jpg", ["A", "B", "X", "D", "E"], NIS="1")], [], e, Q5)[0]
    assert (res.correct, res.wrong, res.grade) == (2, 1, pytest.approx(66.67))


def test_manual_correction_rescores():
    e = exam(corrections={"a.jpg": {"q3": "C", "NIS": "77"}})
    res = ujian.score_results([row("a.jpg", ["A", "B", "", "D", "E"], KodeSoal="101", NIS="")], [], e, Q5)[0]
    assert res.correct == 5 and res.nis == "77" and res.corrected
    assert "No. peserta kosong" not in res.flags


def test_identity_flags_and_student_list():
    students = [
        {"nis": "0012", "nama": "Ani Lestari", "kelas": "9A"},
        {"nis": "0013", "nama": "Budi", "kelas": "9B"},
        {"nis": "0099", "nama": "Tidak Datang", "kelas": "9B"},
    ]
    rows = [
        row("a.jpg", list("ABCDE"), KodeSoal="101", NIS="12", Nama="AMI"),  # leading zeros
        row("b.jpg", list("ABCDE"), KodeSoal="101", NIS="13"),
        row("c.jpg", list("ABCDE"), KodeSoal="101", NIS="13"),  # duplicate
        row("d.jpg", list("ABCDE"), KodeSoal="101", NIS="55"),  # unknown
        row("e.jpg", list("ABCDE"), KodeSoal="101", NIS=""),
    ]
    res = {r.file_id: r for r in ujian.score_results(rows, [{"file_id": "x.jpg"}], exam(), Q5, students)}
    assert res["a.jpg"].name == "Ani Lestari" and res["a.jpg"].kelas == "9A"  # name from the list
    assert any("dobel" in f for f in res["b.jpg"].flags)
    assert "No. peserta tidak ada di daftar peserta" in res["d.jpg"].flags
    assert "No. peserta kosong" in res["e.jpg"].flags
    assert res["x.jpg"].unreadable
    absent = ujian.absent_students(list(res.values()), students)
    assert [s["nis"] for s in absent] == ["0099"]
    # sheets that need a look come first
    ordered = ujian.score_results(rows, [], exam(), Q5, students)
    assert ordered[0].flags and not ordered[-1].flags


def test_sheet_without_nis_matches_students_by_name():
    students = [{"nama": "Andi", "kelas": "A"}, {"nama": "Absen", "kelas": "A"}]
    rows = [row("a.jpg", list("ABCDE"), Nama="ANDI  "), row("b.jpg", list("ABCDE"), Nama="ORANG LAIN")]
    res = {r.file_id: r for r in ujian.score_results(rows, [], ujian.Exam(), Q5, students)}
    assert res["a.jpg"].student["kelas"] == "A" and res["a.jpg"].flags == []  # no "NIS kosong"
    assert res["b.jpg"].flags == ["Nama tidak ada di daftar peserta"]


def test_weighted_questions():
    # q1-q2 worth 2 points, the rest 1 (the general "correct" value)
    e = ujian.Exam(keys={ujian.ALL_CODES: "ABCDE"}, weights={"": "1-2:2"})
    res = ujian.score_results([row("a.jpg", ["A", "B", "X", "D", ""], NIS="1")], [], e, Q5)[0]
    assert (res.correct, res.wrong, res.blank) == (3, 1, 1)
    assert res.score == 2 + 2 + 1 and res.grade == pytest.approx(5 / 7 * 100, abs=0.01)
    e.incorrect = -1  # wrong answers still use the general value
    res = ujian.score_results([row("a.jpg", ["A", "B", "X", "D", ""], NIS="1")], [], e, Q5)[0]
    assert res.score == 4


def test_answers_must_be_in_order():
    # q2 blank: q3-q5 are answered but not counted when the part is in order
    e = ujian.Exam(keys={ujian.ALL_CODES: "ABCDE"}, unmarked=0)
    sheet = row("a.jpg", ["A", "", "C", "D", "X"], NIS="1")
    res = ujian.score_results([sheet], [], e, Q5)[0]
    assert (res.correct, res.wrong, res.blank, res.ignored) == (3, 1, 1, 0) and res.flags == []
    e.in_order = [""]
    res = ujian.score_results([sheet], [], e, Q5)[0]
    assert (res.correct, res.wrong, res.blank, res.ignored) == (1, 0, 1, 3)
    assert res.skipped == ["q3", "q4", "q5"] and res.grade == 20
    assert res.parts[""]["diabaikan"] == 3
    assert res.flags == ["Soal 2 kosong: 3 jawaban sesudahnya tidak dihitung (harus berurutan)"]
    # blanks at the end only: nothing to ignore
    res = ujian.score_results([row("b.jpg", ["A", "B", "C", "", ""], NIS="2")], [], e, Q5)[0]
    assert (res.correct, res.ignored, res.flags) == (3, 0, [])
    # a double mark counts as filled, not as the gap
    res = ujian.score_results([row("c.jpg", ["A", "BC", "C", "D", "E"], NIS="3")], [], e, Q5)[0]
    assert (res.correct, res.wrong, res.ignored) == (4, 1, 0)
    # item analysis: a not-counted answer counts as blank
    analysis = {a["soal"]: a for a in ujian.item_analysis([ujian.score_results([sheet], [], e, Q5)[0]], e, Q5)}
    assert analysis["3"]["Kosong"] == 1 and analysis["3"]["benar"] == 0


def test_in_order_only_for_chosen_parts(tmp_path):
    from openpyxl import load_workbook

    questions = {"I": ["q1", "q2", "q3"], "II": ["q4", "q5"]}
    e = ujian.Exam(keys={ujian.ALL_CODES: {"I": "ABC", "II": "DE"}}, in_order=["I"])
    sheet = row("a.jpg", ["A", "", "C", "", "E"], NIS="1")
    res = ujian.score_results([sheet], [], e, questions)[0]
    assert res.parts["I"]["diabaikan"] == 1 and res.parts["II"]["diabaikan"] == 0
    assert res.correct == 2 and res.skipped == ["q3"]
    assert res.flags == ["Soal I-2 kosong: 1 jawaban sesudahnya tidak dihitung (harus berurutan)"]
    path = ujian.export_excel(tmp_path / "rekap.xlsx", [res], e, questions)
    wb = load_workbook(path)
    header = [c.value for c in wb["Rekap Nilai"][5]]
    assert wb["Rekap Nilai"].cell(6, header.index("Tidak dihitung") + 1).value == 1
    jaw = wb["Jawaban"]
    assert "Kuning" in jaw["A1"].value and jaw.cell(2, 3).value == "I-1"
    assert jaw.cell(3, 5).fill.fgColor.rgb.endswith("FFE699")  # q3 not counted: yellow
    assert jaw.cell(3, 4).fill.fgColor.rgb.endswith("EDEDED")  # q2 blank: grey
    ujian.rename_section_in_exam(e, "I", "Satu", {})
    assert e.in_order == ["Satu"]


def test_switch_template_keeps_keys_per_template():
    e = ujian.Exam(template="/lib/A", code_column="Kode", keys={"1": {"I": "AB"}}, weights={"I": "1:2"},
                   in_order=["I"], corrections={"a.jpg": {"q1": "A"}})
    ujian.switch_template(e, "/lib/A")  # same template: nothing changes
    assert e.keys == {"1": {"I": "AB"}} and e.corrections
    ujian.switch_template(e, "/lib/B")
    assert (e.template, e.code_column, e.keys, e.weights, e.in_order, e.corrections) == ("/lib/B", "", {}, {}, [], {})
    e.keys = {ujian.ALL_CODES: {"X": "C"}}
    ujian.switch_template(e, "/elsewhere/A")  # matched by folder name (moved library)
    assert e.code_column == "Kode" and e.keys == {"1": {"I": "AB"}} and e.weights == {"I": "1:2"}
    assert e.in_order == ["I"] and e.template_keys == {"B": {"code_column": "", "in_order": [], "weights": {},
                                                             "keys": {ujian.ALL_CODES: {"X": "C"}},
                                                             "corrections": {}}}
    assert e.corrections == {"a.jpg": {"q1": "A"}}  # back on A: its corrections are back too


def test_clear_corrections_also_forgets_those_put_aside():
    e = ujian.Exam(template="/lib/A", corrections={"a.jpg": {"q1": "A"}})
    ujian.switch_template(e, "/lib/B")
    e.corrections = {"a.jpg": {"q1": "C"}}
    ujian.clear_corrections(e)  # e.g. another scan folder
    assert e.corrections == {}
    ujian.switch_template(e, "/lib/A")
    assert e.corrections == {}


def test_rekap_state_ignores_keys_of_other_templates(tmp_path):
    (tmp_path / "Results").mkdir()
    (tmp_path / "Results" / "Results_1.csv").write_text("file_id\na.jpg\n", encoding="utf-8")
    e = ujian.Exam(template="/lib/A", output_dir=str(tmp_path), keys={ujian.ALL_CODES: {"": "AB"}})
    state = ujian.rekap_state(e)
    ujian.switch_template(e, "/lib/B")
    ujian.switch_template(e, "/lib/A")  # only looked at another template
    assert e.template_keys and ujian.rekap_state(e) == state
    e.keys[ujian.ALL_CODES][""] = "AC"
    assert ujian.rekap_state(e) != state


def test_forget_folders_forgets_the_old_export(tmp_path):
    e = ujian.Exam(output_dir=str(tmp_path), exported_file=str(tmp_path / "Rekap A.xlsx"), exported_state="x")
    ujian.forget_folders(e, {ujian.PREF_SCANS: False, ujian.PREF_OUTPUT: True})  # same results: still exported
    assert e.exported_file and e.exported_state
    ujian.forget_folders(e, {ujian.PREF_SCANS: False, ujian.PREF_OUTPUT: False})
    assert (e.output_dir, e.exported_file, e.exported_state) == ("", "", "")


def write_results(folder, *fids):
    (folder / "Results").mkdir(parents=True)
    with (folder / "Results" / "Results_1.csv").open("w", newline="", encoding="utf-8") as f:
        f.write("file_id,input_path,output_path,score,q1,q2,q3,q4,q5,NIS\n")
        for i, fid in enumerate(fids):
            f.write(f"{fid},/scan/{fid},/out/{fid},0,A,B,C,D,E,{folder.name}{i}\n")


def test_sheets_with_the_same_name_in_subfolders(tmp_path):
    """Scanners restart their numbering per folder: 7A/scan_0001.jpg and
    7B/scan_0001.jpg are two sheets, a correction is for one of them."""
    out = tmp_path / "hasil"
    write_results(out / "7A", "scan_0001.jpg", "scan_0002.jpg")
    write_results(out / "7B", "scan_0001.jpg")
    write_results(out, "lepas.jpg")
    rows, _ = ujian.read_results(out)
    assert sorted(r["file_id"] for r in rows) == ["7A/scan_0001.jpg", "7A/scan_0002.jpg", "7B/scan_0001.jpg", "lepas.jpg"]
    e = ujian.Exam(keys={ujian.ALL_CODES: "ABCDE"}, corrections={"7B/scan_0001.jpg": {"q1": "E"}})
    res = {r.file_id: r for r in ujian.score_results(rows, [], e, Q5)}
    assert res["7B/scan_0001.jpg"].corrected and res["7B/scan_0001.jpg"].correct == 4
    assert not res["7A/scan_0001.jpg"].corrected and res["7A/scan_0001.jpg"].correct == 5
    assert ujian.done_sheets(out) == [[".", "lepas.jpg"], ["7A", "scan_0001.jpg"], ["7A", "scan_0002.jpg"],
                                      ["7B", "scan_0001.jpg"]]  # fmt: skip

    # saved by v1.3.0: keyed by the bare name; moved when that name is one sheet only
    e = ujian.Exam(keys={ujian.ALL_CODES: "ABCDE"}, corrections={"scan_0002.jpg": {"q1": "E"},
                                                                  "scan_0001.jpg": {"q1": "E"}})  # fmt: skip
    res = {r.file_id: r for r in ujian.score_results(rows, [], e, Q5)}
    assert res["7A/scan_0002.jpg"].corrected and e.corrections["7A/scan_0002.jpg"] == {"q1": "E"}
    assert not res["7A/scan_0001.jpg"].corrected and not res["7B/scan_0001.jpg"].corrected
    assert "scan_0001.jpg" in e.corrections  # ambiguous: left as it was


def test_export_csv(tmp_path):
    import csv

    rows = [row("a.jpg", list("ABCDE"), KodeSoal="101", NIS="1", Nama="ANI"),
            row("b.jpg", ["A", "", "C", "D", "A"], KodeSoal="101", NIS="2", Nama="BUDI")]
    e = exam(kkm=75, in_order=[""])
    results = ujian.score_results(rows, [{"file_id": "rusak.jpg"}], e, Q5)
    path = ujian.export_csv(tmp_path / "rekap.csv", results, e, Q5)
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    with path.open(encoding="utf-8-sig", newline="") as f:
        table = list(csv.reader(f))
    head = table[0]
    assert head[:5] == ["No", "No. Peserta", "Nama", "Kelompok", "Kode Naskah"]
    assert "Tidak dihitung" in head and "Keterangan" in head and head[-5:] == [f"Jawaban {i}" for i in range(1, 6)]
    by_nis = {r[1]: dict(zip(head, r)) for r in table[1:]}
    assert by_nis["1"]["Nilai"] == "100.0" and by_nis["1"]["Keterangan"] == "Lulus"
    assert by_nis["2"]["Tidak dihitung"] == "3" and by_nis["2"]["Jawaban 2"] == "" and by_nis["2"]["Jawaban 5"] == "A"
    assert len(table) == 4 and by_nis[""]["File"] == "rusak.jpg" and by_nis[""]["Nilai"] == ""


def test_parse_weights():
    w, problems = ujian.parse_weights("1-10:2, 11-20:1", 20)
    assert problems == [] and len(w) == 20 and w[1] == 2 and w[20] == 1
    w, problems = ujian.parse_weights("1-3: 2,5; 4=1.5", 5)
    assert problems == [] and w == {1: 2.5, 2: 2.5, 3: 2.5, 4: 1.5}
    _, problems = ujian.parse_weights("1-10:2, 5:3", 20)
    assert problems == ["soal 5 punya dua bobot"]
    _, problems = ujian.parse_weights("1-30:2", 20)
    assert "hanya sampai 20" in problems[0]
    _, problems = ujian.parse_weights("satu:2", 20)
    assert "tidak dikenali" in problems[0]


def test_parse_key():
    assert ujian.parse_key("1.a 2.B\n3c, d-e") == ("ABCDE", [])
    assert ujian.parse_key("ABX") == ("AB", ["X"])


# ---------------------------------------------------------------------------
# Student list, analysis, export
# ---------------------------------------------------------------------------


def test_student_list_from_excel_and_csv(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["No", "NISN", "Nama Siswa", "Kelas"])
    ws.append([1, "0012", "Ani", "9A"])
    ws.append([2, "0013", "Budi", "9B"])
    wb.save(tmp_path / "siswa.xlsx")
    students, headers, cols = ujian.load_students(tmp_path / "siswa.xlsx")
    assert cols == {"nis": "NISN", "nama": "Nama Siswa", "kelas": "Kelas"}
    assert students[0] == {"nis": "0012", "nama": "Ani", "kelas": "9A"}  # zeros kept

    (tmp_path / "siswa.csv").write_text("NIS;Nama;Rombel\n0012;Ani;9A\n", encoding="utf-8")
    students, _, cols = ujian.load_students(tmp_path / "siswa.csv")
    assert cols["kelas"] == "Rombel" and students[0]["nis"] == "0012"


def test_student_list_company_headers():
    assert ujian.guess_student_columns(["No", "NIP", "Nama Karyawan", "Divisi"]) == {
        "nis": "NIP", "nama": "Nama Karyawan", "kelas": "Divisi"}  # fmt: skip
    assert ujian.guess_student_columns(["No. Peserta", "Nama", "Kelompok"]) == {
        "nis": "No. Peserta", "nama": "Nama", "kelas": "Kelompok"}  # fmt: skip


def test_item_analysis():
    rows = [
        row("a.jpg", list("ABCDE"), KodeSoal="101", NIS="1"),
        row("b.jpg", ["A", "C", "", "AB", "E"], KodeSoal="101", NIS="2"),
        row("c.jpg", list("EDCBA"), KodeSoal="102", NIS="3"),
    ]
    results = ujian.score_results(rows, [], exam(), Q5)
    items = {a["soal"]: a for a in ujian.item_analysis(results, exam(), Q5)}
    assert items["1"]["benar"] == 100.0 and items["1"]["kategori"] == "Mudah"
    assert items["2"]["benar"] == pytest.approx(66.7)
    assert (items["3"]["Kosong"], items["4"]["Ganda"]) == (1, 1)
    assert items["1"]["kunci"] == "101: A, 102: E"


def test_export_excel(tmp_path):
    from openpyxl import load_workbook

    students = [{"nis": "1", "nama": "Ani", "kelas": "9A"}, {"nis": "9", "nama": "Absen", "kelas": "9A"}]
    rows = [
        row("a.jpg", list("ABCDE"), KodeSoal="101", NIS="1"),
        row("b.jpg", ["A", "", "C", "D", "A"], KodeSoal="101", NIS="2"),
    ]
    e = exam(kkm=75, name="PAS IPA Kelas 9")
    results = ujian.score_results(rows, [], e, Q5, students)
    path = ujian.export_excel(tmp_path / "rekap.xlsx", results, e, Q5, students)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Rekap Nilai", "Jawaban", "Analisis Butir", "Tidak Ikut Ujian"]
    ws = wb["Rekap Nilai"]
    assert ws["A1"].value == "PAS IPA Kelas 9"
    header = [c.value for c in ws[5]]
    assert header[:10] == ["No", "No. Peserta", "Nama", "Kelompok", "Kode Naskah", "Benar", "Salah", "Kosong", "Skor", "Nilai"]
    data = {ws.cell(r, 2).value: [c.value for c in ws[r]] for r in range(6, 8)}
    assert header[10:12] == ["Peringkat", "Keterangan"]
    assert data["1"][9:12] == [100, 1, "Lulus"]
    assert data["2"][9:12] == [60, 2, "Tidak lulus"]
    flagged_row = next(r for r in range(6, 8) if ws.cell(r, 2).value == "2")
    assert ws.cell(flagged_row, 1).fill.fgColor.rgb.endswith("F8CBAD")  # not in list -> red
    assert [c.value for c in wb["Tidak Ikut Ujian"][2]] == [1, "9", "Absen", "9A"]
    assert [c.value for c in wb["Analisis Butir"][1]][:3] == ["Soal", "Kunci", "Bobot"]
    jaw = wb["Jawaban"]
    r2 = next(r for r in range(2, 4) if jaw.cell(r, 1).value == "2")
    assert jaw.cell(r2, 5).fill.fgColor.rgb.endswith("EDEDED")  # blank q2: grey
    assert jaw.cell(r2, 8).fill.fgColor.rgb.endswith("F8CBAD")  # wrong q5: red


def test_rank_and_rekap_order():
    rows = [
        row("c.jpg", list("ABCDE"), KodeSoal="101", NIS="3", Nama="CICI"),
        row("a.jpg", ["A", "", "C", "D", "A"], KodeSoal="101", NIS="1", Nama="ANI"),
        row("b.jpg", list("ABCDE"), KodeSoal="101", NIS="2", Nama="BUDI"),
    ]
    e = exam()
    results = ujian.score_results(rows, [{"file_id": "x.jpg"}], e, Q5)
    by_grade = ujian.order_results(results, "nilai")
    assert [(r.file_id, r.rank) for r in by_grade] == [("b.jpg", 1), ("c.jpg", 1), ("a.jpg", 3), ("x.jpg", None)]
    assert [r.file_id for r in ujian.order_results(results, "file")] == ["a.jpg", "b.jpg", "c.jpg", "x.jpg"]
    assert [r.name for r in ujian.order_results(results, "nama")][:3] == ["ANI", "BUDI", "CICI"]


def test_participant_reports_pdf(tmp_path):
    fitz = pytest.importorskip("fitz")
    rows = [row("a.jpg", ["A", "", "C", "D", "A"], KodeSoal="101", NIS="1", Nama="ANI")]
    e = exam(kkm=50, name="Tes Rekrutmen")
    results = ujian.score_results(rows, [{"file_id": "rusak.jpg"}], e, Q5)
    path = ujian.export_reports_pdf(tmp_path / "laporan.pdf", results, e, Q5, title="Tes Rekrutmen")
    with fitz.open(str(path)) as doc:
        assert doc.page_count == 1  # unreadable sheets get no page
        text = doc[0].get_text()
    for expected in ("Tes Rekrutmen", "ANI", "No. Peserta", "Nilai", "Peringkat", "Lulus"):
        assert expected in text, expected


def test_template_file_roundtrip(tmp_path):
    src = tmp_path / "lib" / "LJK A"
    src.mkdir(parents=True)
    (src / "template.json").write_text(json.dumps({"fieldBlocks": {}, "bubbleDimensions": [10, 10],
                                                   "pageDimensions": [100, 100]}), encoding="utf-8")  # fmt: skip
    (src / ujian.META_FILE).write_text(json.dumps({"name": "LJK A"}), encoding="utf-8")
    packed = ujian.export_template_file(src, tmp_path / "LJK A.ljk")
    other = tmp_path / "lib2"
    folder = ujian.import_template_file(packed, other)
    assert folder == other / "LJK A" and (folder / "template.json").is_file()
    assert ujian.import_template_file(packed, other) == other / "LJK A (2)"  # never overwrites

    import zipfile

    bad = tmp_path / "jahat.ljk"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("template.json", "{}")
        z.writestr("../keluar.txt", "x")
    with pytest.raises(ValueError, match="tidak wajar"):
        ujian.import_template_file(bad, other)
    (tmp_path / "bukan.ljk").write_text("halo")
    with pytest.raises(ValueError, match="Bukan file template"):
        ujian.import_template_file(tmp_path / "bukan.ljk", other)


def test_exam_settings_roundtrip(tmp_path):
    e = exam(scans_dir="C:/ujian", corrections={"a.jpg": {"q1": "B"}})
    e.save(tmp_path / "s.json")
    assert ujian.Exam.load(tmp_path / "s.json") == e
    assert ujian.Exam.load(tmp_path / "missing.json") == ujian.Exam()


# ---------------------------------------------------------------------------
# Template library + wizard processing on the real LJK
# ---------------------------------------------------------------------------

PDF = Path(__file__).parent / "data" / "ljk_kosong_smpn3.pdf"


@pytest.fixture(scope="module")
def real_template():
    sheet = tb.prepare_sheet(tb.load_image(PDF), source_name="LJK_kosong.pdf")
    _, blocks = tb.detect_blocks(sheet.page)
    template, problems = tb.build_template(blocks, sheet)
    assert not problems
    return sheet, blocks, template


@needs("ljk_kosong_smpn3.pdf")
def test_library_save_list_edit_import_delete(tmp_path, real_template):
    sheet, blocks, template = real_template
    lib = tmp_path / "lib"
    folder = ujian.save_to_library("LJK SMPN 3", sheet, template, blocks, lib)
    again = ujian.save_to_library("LJK SMPN 3", sheet, template, blocks, lib)
    assert again.name == "LJK SMPN 3 (2)"
    items = ujian.list_templates(lib)
    assert [t.name for t in items] == ["LJK SMPN 3", "LJK SMPN 3"]
    t = items[0]
    assert len(t.questions) == 100 and set(t.columns) == {"Nama", "NIS", "KodeSoal", "TglLahir"}
    assert t.preview is not None and ujian.can_edit(folder)

    sheet2, blocks2, name = ujian.load_for_editing(folder)
    assert name == "LJK SMPN 3" and sheet2.uses_markers and sheet2.source_name == "LJK_kosong.pdf"
    assert [(b.rows, b.cols, b.group, b.start) for b in blocks2] == [
        (b.rows, b.cols, b.group, b.start) for b in blocks
    ]
    t2, problems = tb.build_template(blocks2, sheet2)
    assert not problems and t2 == template  # editing round-trips exactly

    old = tmp_path / "lama"
    tb.save_template_folder(old, sheet, template)
    imported = ujian.import_template_folder(old, "Dari mode lanjutan", lib)
    assert ujian.read_library_template(imported).name == "Dari mode lanjutan"
    assert not ujian.can_edit(imported)

    ujian.delete_template(again)
    assert len(ujian.list_templates(lib)) == 2


@needs("ljk_kosong_smpn3.pdf")
def test_wizard_processing_two_codes(tmp_path, real_template):
    """Template from the library, scans in their own folder (left untouched),
    two exam codes with different keys, a student list, an unreadable page
    and the blank sheet lying in the scan folder."""
    import test_real_ljk as real

    sheet, blocks, template = real_template
    lib_folder = ujian.save_to_library("LJK", sheet, template, blocks, tmp_path / "lib")
    fields, answers, _ = real.ground_truth()
    rnd = random.Random(3)
    keys = {"101": "".join(rnd.choice("ABCDE") for _ in range(100)),
            "102": "".join(rnd.choice("ABCDE") for _ in range(100))}  # fmt: skip

    scans = tmp_path / "scan kelas 9"
    scans.mkdir()
    shutil.copy(PDF, scans / "LJK_kosong.pdf")
    blank = 255 * __import__("numpy").ones((2339, 1654, 3), "uint8")  # page without markers
    cv2.imencode(".jpg", blank)[1].tofile(str(scans / "halaman_kosong.jpg"))
    people = [("ANI", "00000000012", "101", {}), ("BUDI", "00000000013", "102", {"angle": 1.0, "jpeg": 85})]
    expected = {}
    for k, (nama, nis, kode, params) in enumerate(people):
        resp = {"Nama": nama, "NIS": nis, "KodeSoal": kode, "TglLahir": "010110"}
        for q in range(1, 101):
            resp[f"q{q}"] = rnd.choice("ABCDE")
        img = synth.distort(real.render_filled(fields, answers, resp, nama), seed=k, **params)
        name = f"siswa_{k + 1}.jpg"
        cv2.imencode(".jpg", img)[1].tofile(str(scans / name))
        right = sum(resp[f"q{q}"] == keys[kode][q - 1] for q in range(1, 101))
        expected[name] = (nis, kode, right)
    before = sorted(p.name for p in scans.iterdir())

    out = tmp_path / "hasil"
    proc = subprocess.run(
        [sys.executable, str(ROOT / "gui.py"), "--process", str(scans), str(out), "--template", str(lib_folder)],
        capture_output=True, text=True, timeout=600,
    )
    log = (out / "omrchecker_log.txt").read_text(encoding="utf-8") if out.exists() else ""
    assert proc.returncode == 0, log[-3000:] + proc.stderr[-2000:]
    assert sorted(p.name for p in scans.iterdir()) == before  # nothing written to the scans

    rows, errors = ujian.read_results(out)
    # the blank sheet is read too (nothing marked) and hidden when scoring
    assert sorted(r["file_id"] for r in rows) == ["LJK_kosong.png", "siswa_1.jpg", "siswa_2.jpg"]
    assert [e["file_id"] for e in errors] == ["halaman_kosong.jpg"]

    e = ujian.Exam(code_column="KodeSoal", keys=keys)
    students = [{"nis": "12", "nama": "Ani Lestari", "kelas": "9A"},
                {"nis": "13", "nama": "Budi Santoso", "kelas": "9A"}]  # fmt: skip
    questions = ujian.read_library_template(lib_folder).questions
    res = {r.file_id: r for r in ujian.score_results(rows, errors, e, questions, students,
                                                     blank_source="LJK_kosong.pdf")}  # fmt: skip
    assert "LJK_kosong.png" not in res
    for name, (nis, kode, right) in expected.items():
        r = res[name]
        assert (r.nis, r.code, r.correct) == (nis, kode, right), name
        assert r.grade == right and r.flags == [] and r.student is not None
    assert res["halaman_kosong.jpg"].unreadable
    path = ujian.export_excel(tmp_path / "rekap.xlsx", list(res.values()), e, questions, students)
    assert path.exists()


@needs("ljk_kosong_smpn3.pdf")
def test_rename_section_moves_labels_keys_and_corrections(tmp_path, real_template):
    import test_medion_ljk as medion

    sheet = tb.prepare_sheet(tb.load_image(medion.SCAN))
    bubbles, blocks = tb.detect_blocks(sheet)
    blocks = medion.label_like_a_teacher(bubbles, blocks)
    template, _ = tb.build_template(blocks, sheet)
    folder = ujian.save_to_library("Medion", sheet, template, blocks, tmp_path / "lib")
    assert ujian.section_thumbnail(folder, "TYPE") is not None
    assert ujian.section_box(folder, "I")[2] > ujian.section_box(folder, "TYPE")[2] - 1000

    mapping = ujian.rename_section(folder, "Motivasi", "Minat")
    assert mapping == {f"Motivasi_{i}": f"Minat_{i}" for i in range(1, 11)}
    info = ujian.read_library_template(folder)
    assert list(info.sections) == ["I", "II", "III", "Minat", "TYPE"]
    assert len(info.sections["Minat"]) == 10
    with pytest.raises(ValueError):
        ujian.rename_section(folder, "Minat", "I")  # name taken

    e = ujian.Exam(keys={ujian.ALL_CODES: {"Motivasi": "AAAA"}}, weights={"Motivasi": "1-2:2"},
                   corrections={"a.jpg": {"Motivasi_3": "B", "Nama": "X"}})  # fmt: skip
    ujian.rename_section_in_exam(e, "Motivasi", "Minat", mapping)
    assert e.keys[ujian.ALL_CODES] == {"Minat": "AAAA"} and e.weights == {"Minat": "1-2:2"}
    assert e.corrections == {"a.jpg": {"Minat_3": "B", "Nama": "X"}}
