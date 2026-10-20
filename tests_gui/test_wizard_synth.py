"""The wizard end to end on a synthetic LJK (no private sample sheets needed):
step 5 must wait for the "Proses LJK" click, then read every sheet; an empty
page is flagged without stopping the batch.

Needs Tk and a display (skipped otherwise); headless Linux: xvfb-run.
"""

import sys
import time
from pathlib import Path

import cv2
import numpy as np
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(Path(__file__).parent)]

tk = pytest.importorskip("tkinter", exc_type=ImportError)

import synth_ljk as synth  # noqa: E402
import template_builder as tb  # noqa: E402
import ujian  # noqa: E402
from test_template_builder import label_like_user  # noqa: E402


def iter_widgets(w):
    for c in w.winfo_children():
        yield c
        yield from iter_widgets(c)


def pump(w, n=10):
    for _ in range(n):
        w.update()


@pytest.fixture
def setup(tmp_path, monkeypatch):
    try:
        tk.Tk().destroy()
    except tk.TclError as e:
        pytest.skip(f"no display: {e}")
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    img, layout = synth.make_blank()
    sheet = tb.prepare_sheet(img)
    _, blocks = tb.detect_blocks(sheet.page)
    blocks = label_like_user([*blocks], layout)
    template, _ = tb.build_template(blocks, sheet)
    lib = tmp_path / "lib"
    ujian.save_to_library("LJK Sintetis", sheet, template, blocks, lib)
    scans = tmp_path / "scan"
    scans.mkdir()
    for k in range(2):
        filled = synth.fill(img, layout, synth.random_responses(layout, seed=k), seed=k)
        cv2.imencode(".jpg", filled)[1].tofile(str(scans / f"siswa_{k + 1}.jpg"))
    # an empty page (e.g. the back of a duplex scan) must not stop the batch
    cv2.imencode(".jpg", np.full((3508, 2480, 3), 235, np.uint8))[1].tofile(str(scans / "kosong.jpg"))
    return lib, scans, tmp_path


def test_process_step_waits_for_click(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    errors = []
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: errors.append(a))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: False)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        assert w.stepper.current == 0
        w.next()
        step = w.steps[1]
        for part, text in step.texts.items():
            count = len(w.template.sections[part])
            text.insert("1.0", "A" * count)
            step.on_key_edit(part)
        w.next()
        w.steps[2].folder.set(str(scans))
        w.next()
        w.steps[3].name.set("Uji")
        w.next()
        assert w.current == 4, errors
        assert w.stepper.current == 4 and w.stepper.reached == 4

        step = w.steps[4]
        pump(w, 30)
        time.sleep(0.5)
        pump(w, 10)
        assert step.proc is None and not step.results, "processing must not start by itself"
        assert "Klik" in step.summary.get() and "3 LJK" in step.summary.get()
        assert step.progress_text.get() == "3 LJK akan diproses"  # the count shows before the click

        step.run_btn.invoke()
        assert step.proc is not None
        deadline = time.time() + 300
        while step.proc is not None:
            w.update()
            time.sleep(0.05)
            assert time.time() < deadline, "processing did not finish"
        pump(w, 20)
        assert not errors, errors
        assert len(step.results) == 3
        bad = [r for r in step.results if r.flags]
        assert [r.file_id for r in bad] == ["kosong.jpg"]
        assert "tidak terbaca" in bad[0].flags[0]
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def alive(w):
    try:
        return bool(w.winfo_exists())
    except tk.TclError:  # the root window itself was destroyed
        return False


def run_to_results(w):
    """Steps 1-5 with all-A keys; returns once step 5 has read every sheet."""
    pump(w)
    w.next()
    step = w.steps[1]
    for part, text in step.texts.items():
        text.insert("1.0", "A" * len(w.template.sections[part]))
        step.on_key_edit(part)
    w.next()
    w.next()
    w.next()
    step = w.steps[4]
    step.run_btn.invoke()
    deadline = time.time() + 300
    while step.proc is not None:
        w.update()
        time.sleep(0.05)
        assert time.time() < deadline
    pump(w, 20)
    return step


def test_close_offers_export_until_rekap_is_current(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)  # confirm "Ujian baru"
    asked = []
    answers = []

    def ask(master, heading, text, buttons, detail=""):
        asked.append((heading, [k for k, *_ in buttons]))
        return answers.pop(0)

    monkeypatch.setattr(wz.ChoiceDialog, "ask", staticmethod(ask))
    out = tmp / "rekap.xlsx"
    monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(out))
    exam = ujian.Exam(scans_dir=str(scans), output_dir=str(tmp / "hasil"))
    exam.save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        run_to_results(w)
        assert w.export_pending()

        # X button: nothing exported yet -> asked; "Batal" keeps the window
        answers.append("batal")
        w.on_close()
        assert asked[-1] == ("Rekap nilai belum diekspor", ["batal", "tutup", "ekspor"])
        assert alive(w)

        # step 6, "Selesai": export from the prompt, then the app closes
        w.next()
        assert w.current == 5
        assert w.steps[5].export_info.get() == "Belum diekspor"
        answers.append("ekspor")
        w.next()  # Selesai
        assert out.exists()
        assert not alive(w)
        saved = ujian.Exam.load(tmp / "sesi.json")
        assert saved.exported_file == str(out) and saved.exported_state == ujian.rekap_state(saved)
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass

    # reopen with both folders remembered: rekap is current -> "Selesai" shows
    # the closing screen, no export prompt
    ujian.save_prefs({ujian.PREF_SCANS: True, ujian.PREF_OUTPUT: True}, tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        for _ in range(5):
            w.next()
        assert w.current == 5 and not w.export_pending()
        assert w.steps[5].export_info.get() == "✓ Sudah diekspor: rekap.xlsx"
        answers.append("folder")
        opened = []
        monkeypatch.setattr(wz, "open_path", opened.append)
        w.next()
        assert asked[-1] == ("Pemeriksaan selesai", ["buka", "folder", "baru", "tutup"])
        assert opened == [w.exam.output_dir] and alive(w)

        # a correction after the export makes the rekap out of date again
        w.exam.corrections = {"siswa_1.jpg": {"q1": "B"}}
        assert w.export_pending()
        answers.append("tutup")
        w.on_close()
        assert asked[-1][0] == "Rekap belum diperbarui"
        assert not alive(w)
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_choice_dialog_returns_clicked_key(setup):
    import wizard as wz

    lib, _, tmp = setup
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        d = wz.ChoiceDialog(w, "Judul", "Teks", [("a", "Satu", False), ("b", "Dua", True)], detail="rincian")
        pump(w)
        assert str(d.buttons["b"]["style"]) == "Accent.TButton"
        d.buttons["b"].invoke()
        assert d.choice == "b" and not d.winfo_exists()
        d = wz.ChoiceDialog(w, "Judul", "Teks", [("a", "Satu", False)])
        d.event_generate("<Escape>")
        pump(w)
        assert d.choice is None
    finally:
        w.destroy()


def test_old_exam_can_be_reopened_from_its_output_folder(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)  # confirm "Ujian baru"
    monkeypatch.setattr(wz.ChoiceDialog, "ask", staticmethod(lambda *a, **k: "tutup"))  # skip the export
    hasil = tmp / "hasil"
    ujian.Exam(scans_dir=str(scans), output_dir=str(hasil), name="Tes Rekrutmen").save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        run_to_results(w)
        w.exam.corrections = {"siswa_1.jpg": {"q1": "B"}}
        w.exam.save(w.settings)
        assert (hasil / ujian.EXAM_FILE).is_file()
        w.next()
        w.steps[5].new_exam()  # the session now holds an empty exam
        assert w.exam.corrections == {} and w.exam.output_dir == ""

        monkeypatch.setattr(wz.filedialog, "askdirectory", lambda **k: str(hasil))
        w.open_exam()
        pump(w)
        assert w.current == 5
        assert w.exam.name == "Tes Rekrutmen" and w.exam.corrections == {"siswa_1.jpg": {"q1": "B"}}
        assert all(step.exam is w.exam for step in w.steps)
        assert len(w.steps[5].results) == 3
        assert ujian.Exam.load(tmp / "sesi.json").output_dir == str(hasil)  # it is the current exam again
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_open_exam_messages(setup, tmp_path):
    lib, _, _ = setup
    with pytest.raises(ValueError, match="bukan folder hasil"):
        ujian.open_exam(tmp_path, lib)
    folder = tmp_path / "hasil_lain"
    folder.mkdir()
    # made on another computer: template path differs, same name in this library
    exam = ujian.Exam(template=str(tmp_path / "pc-lain" / "LJK Sintetis"), output_dir=str(folder))
    exam.save(tmp_path / "sesi.json")
    opened = ujian.open_exam(folder, lib)
    assert Path(opened.template) == lib / "LJK Sintetis"
    exam.template = str(tmp_path / "Tidak Ada")
    exam.save(tmp_path / "sesi.json")
    with pytest.raises(ValueError, match="Tidak Ada"):
        ujian.open_exam(folder, lib)


def test_about_shows_license(setup, monkeypatch):
    import wizard as wz

    lib, _, tmp = setup
    shown = []
    answers = ["lisensi", "pihak"]
    monkeypatch.setattr(wz.ChoiceDialog, "ask",
                        staticmethod(lambda m, h, t, b, detail="": (shown.append((h, t)), answers.pop(0))[1]))  # fmt: skip
    monkeypatch.setattr(wz, "show_text", lambda parent, title, text: shown.append(text))
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.about()
        heading, text = shown[0]
        assert heading == f"Pemindai LJK v{wz.VERSION}" and "AGPL-3.0" in text
        assert shown[1] == (ROOT / "LICENSE").read_text(encoding="utf-8")
        w.about()
        assert "MIT License" in shown[3] and "Udayraj" in shown[3]  # upstream notice
    finally:
        w.destroy()


def test_stopped_run_resumes_without_rereading(setup, monkeypatch):
    import csv

    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    asked = []
    monkeypatch.setattr(wz.ChoiceDialog, "ask", staticmethod(lambda m, h, t, b, detail="": (asked.append(t), "lanjut")[1]))
    hasil = tmp / "hasil"
    ujian.Exam(scans_dir=str(scans), output_dir=str(hasil)).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        step = run_to_results(w)
        assert step.done_count == 3 and step.progress["value"] == 3  # progress from the worker, not the log
        # as if the run had been stopped before siswa_2.jpg was read
        results_csv = next(hasil.rglob("Results/Results_*.csv"))
        with open(results_csv, newline="", encoding="utf-8") as fh:
            rows = list(csv.reader(fh))
        with open(results_csv, "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh, quoting=csv.QUOTE_NONNUMERIC).writerows(r for r in rows if r[0] != "siswa_2.jpg")
        assert len(ujian.done_sheets(hasil)) == 2

        step.run_btn.invoke()
        assert asked == ["2 dari 3 LJK sudah diperiksa sebelumnya."]
        deadline = time.time() + 300
        while step.proc is not None:
            w.update()
            time.sleep(0.05)
            assert time.time() < deadline
        pump(w, 20)
        assert "".join(step.log).count("Opening image") == 1  # only the missing sheet was read
        assert len(step.results) == 3
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_rekap_order_and_participant_reports(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: False)
    names = tmp / "peserta.csv"  # one listed twice, one without a sheet
    names.write_text("No. Peserta,Nama\n1,Ani\n1,Ani lagi\n99,Tidak Datang\n", encoding="utf-8")
    ujian.Exam(scans_dir=str(scans), output_dir=str(tmp / "hasil"), students_file=str(names),
               students_cols={"nis": "No. Peserta", "nama": "Nama"}).save(tmp / "sesi.json")  # fmt: skip
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        run_to_results(w)
        assert "No. Peserta dobel di daftar: 1" in w.steps[2].student_info.get()
        w.next()
        step = w.steps[5]
        assert "tidak ikut: " in step.stats.get() and step.absent_row.winfo_manager() == "pack"
        dialog = step.show_absent()
        pump(w)
        tree = next(c for c in iter_widgets(dialog) if isinstance(c, wz.ttk.Treeview))
        assert "99" in [str(tree.item(i)["values"][0]) for i in tree.get_children()]
        dialog.destroy()
        assert step.tree.get_children() and step.analysis[0]["terbanyak"]
        step.order.set(ujian.RANK_ORDERS["nilai"])
        step.on_order()
        assert w.exam.rekap_order == "nilai"
        pdf = tmp / "laporan.pdf"
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(pdf))
        monkeypatch.setattr(wz.ChoiceDialog, "ask", classmethod(lambda cls, *a, **k: "satu"))
        step.export_reports()
        import fitz

        with fitz.open(str(pdf)) as doc:
            assert doc.page_count == 2  # the empty page has no report
        monkeypatch.setattr(wz.ChoiceDialog, "ask", classmethod(lambda cls, *a, **k: "pisah"))
        monkeypatch.setattr(wz.filedialog, "askdirectory", lambda **k: str(tmp / "kirim"))
        step.export_reports()
        assert len(list((tmp / "kirim").glob("*/*.pdf"))) == 2
        csv_path = tmp / "rekap.csv"
        monkeypatch.setattr(wz.messagebox, "showinfo", lambda *a, **k: None)
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(csv_path))
        assert step.export_csv()
        import csv

        with csv_path.open(encoding="utf-8-sig", newline="") as f:
            table = list(csv.reader(f))
        assert table[0][:3] == ["No", "No. Peserta", "Nama"] and "Nilai" in table[0]
        assert table[0][-1].startswith("Jawaban ") and len(table) == 1 + len(step.results)
        assert w.exam.exported_file == str(csv_path) and not w.export_pending()
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_advanced_export(setup, monkeypatch):
    import wizard as wz
    from openpyxl import load_workbook

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: False)
    ujian.Exam(scans_dir=str(scans), output_dir=str(tmp / "hasil")).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        run_to_results(w)
        w.next()
        step = w.steps[5]
        out = tmp / "rekap-sendiri.xlsx"
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(out))
        monkeypatch.setattr(wz.simpledialog, "askstring", lambda *a, **k: "Format sekolah")
        seen = {}

        def drive(d):
            pump(d, 3)
            items = {d.rows[i][1]: i for i in d.tree.get_children()}
            assert all(d.rows[i][0] for i in items.values())  # first time: every column
            d.check_all(False)
            for name in ("Nilai", "Nama", "No. Peserta"):
                d.tree.selection_set(items[name])
                d.toggle()
            d.tree.selection_set(items["No. Peserta"])
            d.on_select()
            d.title_var.set("NIS")
            d.move(-1)
            d.move(-1)  # to the top
            d.tree.selection_set(items["Nilai"])
            d.move(1)  # past Nama
            d.save_format()
            seen["values"] = list(d.fmt_box.cget("values"))
            d.ok()

        monkeypatch.setattr(step, "wait_window", drive)
        assert step.export_advanced() == str(out)
        assert seen["values"] == ["Format sekolah"]
        table = [[c.value for c in r] for r in load_workbook(out)["Rekap"].iter_rows()]
        assert table[0] == ["NIS", "Nama", "Nilai"] and len(table) == 1 + len(step.results)
        assert w.exam.exported_file == str(out) and not w.export_pending()
        formats = ujian.load_export_formats(tmp / "sesi.json")
        assert formats["Format sekolah"]["kolom"] == [["No. Peserta", "NIS"], ["Nama", "Nama"], ["Nilai", "Nilai"]]

        def again(d):  # the last settings come back; CSV this time
            pump(d, 3)
            assert d.columns() == [["No. Peserta", "NIS"], ["Nama", "Nama"], ["Nilai", "Nilai"]]
            d.kind.set("csv")
            d.ok()

        csv_out = tmp / "rekap-sendiri.csv"
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(csv_out))
        monkeypatch.setattr(step, "wait_window", again)
        assert step.export_advanced() == str(csv_out)
        assert csv_out.read_text(encoding="utf-8-sig").splitlines()[0] == "NIS,Nama,Nilai"
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_template_export_and_import_buttons(setup, monkeypatch):
    import wizard as wz

    lib, _, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "showinfo", lambda *a, **k: None)
    packed = tmp / "LJK Sintetis.ljk"
    monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(packed))
    monkeypatch.setattr(wz.filedialog, "askopenfilename", lambda **k: str(packed))
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        step = w.steps[0]
        step.export_file()
        assert packed.is_file()
        step.import_file()
        assert [t.name for t in step.templates] == ["LJK Sintetis", "LJK Sintetis"]
        assert step.selected().path == lib / "LJK Sintetis (2)"
    finally:
        w.destroy()


def test_stop_reports_how_far_it_got(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)  # no answer key: read only
    ujian.Exam(scans_dir=str(scans), output_dir=str(tmp / "hasil")).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.next()
        w.next()
        w.next()
        w.next()
        step = w.steps[4]
        step.run_btn.invoke()
        pump(w, 5)
        step.stop_btn.invoke()
        pump(w, 5)
        assert step.proc is None and step.progress_text.get() == "Dihentikan"
        assert step.summary.get().startswith("Dihentikan:") and "dari 3 LJK" in step.summary.get()
    finally:
        w.destroy()


def test_full_rerun_keeps_exam_file(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.ChoiceDialog, "ask", staticmethod(lambda *a, **k: "ulang"))
    hasil = tmp / "hasil"
    ujian.Exam(scans_dir=str(scans), output_dir=str(hasil)).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        step = run_to_results(w)
        (scans / "baru.jpg").write_bytes((scans / "siswa_1.jpg").read_bytes())  # a 4th scan arrives
        (hasil / "catatan saya.txt").write_text("jangan dihapus")
        w.go(4)
        step.run_btn.invoke()
        deadline = time.time() + 300
        while step.proc is not None:
            w.update()
            time.sleep(0.05)
            assert time.time() < deadline
        pump(w, 20)
        assert "".join(step.log).count("Opening image") == 4  # everything read again
        assert (hasil / ujian.EXAM_FILE).is_file() and (hasil / "catatan saya.txt").is_file()
        assert len(step.results) == 4
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_remember_folder_checkboxes(setup, monkeypatch):
    """Steps 3 and 4: 'Ingat folder ini' is off by default; after a finished
    (exported) exam or 'Ujian baru' only the checked folders are kept."""
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    hasil = tmp / "hasil"
    done = ujian.Exam(scans_dir=str(scans), output_dir=str(hasil), exported_file=str(tmp / "rekap.xlsx"),
                      corrections={"siswa_1.jpg": {"q1": "B"}})  # fmt: skip
    done.save(tmp / "sesi.json")

    # default: nothing remembered, the next exam starts with empty folders
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        assert not w.steps[2].remember.var.get() and not w.steps[3].remember.var.get()
        assert w.exam.scans_dir == "" and w.exam.output_dir == "" and w.exam.corrections == {}
        w.next()
        w.next()
        assert w.current == 2 and w.steps[2].folder.get() == ""
        # check "remember" in step 3: saved right away as a preference
        w.steps[2].remember.invoke()
        assert ujian.load_prefs(tmp / "sesi.json") == {ujian.PREF_SCANS: True, ujian.PREF_OUTPUT: False,
                                                     ujian.PREF_DASHBOARD: False}
        w.steps[2].folder.set(str(scans))
        w.next()
        assert w.current == 3 and w.steps[3].folder.get().endswith("scan_hasil")

        # "Ujian baru": the remembered scan folder stays, the output folder is cleared
        w.exam.output_dir = str(hasil)
        w.steps[5].new_exam()
        assert w.exam.scans_dir == str(scans) and w.exam.output_dir == ""
    finally:
        w.destroy()

    # both remembered: a finished exam keeps its folders on the next start
    ujian.save_prefs({ujian.PREF_SCANS: True, ujian.PREF_OUTPUT: True}, tmp / "sesi.json")
    done.save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        assert w.steps[2].remember.var.get() and w.steps[3].remember.var.get()
        assert w.exam.scans_dir == str(scans) and w.exam.output_dir == str(hasil)
    finally:
        w.destroy()


def test_archive_old_scans(setup, monkeypatch):
    """Step 3 'Arsipkan scan lama': the scans move to _arsip/<date time> (not
    deleted, not read again), corrections are dropped and step 4 then
    replaces the old results, so a re-scan does not double up."""
    import wizard as wz

    lib, scans, tmp = setup
    asked = []
    monkeypatch.setattr(wz.ChoiceDialog, "ask", classmethod(lambda cls, *a, **k: asked.append(a) or "arsip"))
    hasil = tmp / "hasil"
    (hasil / "Results").mkdir(parents=True)
    (hasil / "Results" / "Results_01PM.csv").write_text("file_id\n")
    ujian.Exam(scans_dir=str(scans), output_dir=str(hasil), corrections={"siswa_1.jpg": {"q1": "B"}}).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.next()
        for part, text in w.steps[1].texts.items():
            text.insert("1.0", "A" * len(w.template.sections[part]))
            w.steps[1].on_key_edit(part)
        w.next()
        assert w.current == 2
        step = w.steps[2]
        assert step.count.get().startswith("3 file") and str(step.archive_btn["state"]) == "normal"
        step.archive_btn.invoke()
        assert "3 file scan" in asked[0][2]
        archived = list((scans / wz.ARCHIVE_DIR).rglob("*.jpg"))
        assert len(archived) == 3 and not list(scans.glob("*.jpg"))
        assert wz.scan_files(scans) == [] and step.count.get().startswith("0 file")
        assert str(step.archive_btn["state"]) == "disabled" and w.exam.corrections == {}

        # the re-scan lands in the same folder; old results get replaced in step 4
        (scans / "scan_0001.jpg").write_bytes(archived[0].read_bytes())
        step.scan()
        w.steps[3].policy.set("keep")
        w.next()
        assert w.current == 3 and w.steps[3].policy.get() == "replace"

        # only the new scan is read: the archived ones are skipped by the run too
        w.next()
        step = w.steps[4]
        assert w.current == 4 and step.progress_text.get() == "1 LJK akan diproses"
        step.run_btn.invoke()
        deadline = time.time() + 300
        while step.proc is not None:
            w.update()
            time.sleep(0.05)
            assert time.time() < deadline, "processing did not finish"
        pump(w, 20)
        assert [r.file_id for r in step.results] == ["scan_0001.jpg"]
        assert not list(hasil.rglob(wz.ARCHIVE_DIR))
    finally:
        w.destroy()


def test_unfinished_exam_keeps_folders(setup):
    """An exam not exported yet (e.g. processing stopped halfway) is resumed
    with its folders, whatever the checkboxes say."""
    import wizard as wz

    lib, scans, tmp = setup
    ujian.Exam(scans_dir=str(scans), output_dir=str(tmp / "hasil")).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        assert w.exam.scans_dir == str(scans) and w.exam.output_dir == str(tmp / "hasil")
    finally:
        w.destroy()


def test_key_step_drops_keys_of_other_templates(setup, monkeypatch):
    """Keys saved for sections the active template does not have (from an
    earlier, since deleted template) used to block step 2 with "bagian ini
    hanya 0 soal" although nothing on screen could fix it."""
    import wizard as wz

    lib, scans, tmp = setup
    errors = []
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: errors.append(a[1]))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.exam.keys = {ujian.ALL_CODES: {"Motivasi1": "ABC", "Motivasi2": "ABCDABC"}}
        w.exam.weights = {"Motivasi1": "1:2"}
        w.next()
        step = w.steps[1]
        assert "Motivasi1" not in step.texts
        part = next(iter(step.texts))
        step.texts[part].insert("1.0", "A")
        step.on_key_edit(part)
        w.next()
        assert w.current == 2, errors
        assert w.exam.keys[ujian.ALL_CODES] == {part: "A"}
        assert w.exam.weights == {}
    finally:
        w.destroy()


def test_key_step_in_order_checkbox(setup, monkeypatch):
    """Step 2: 'Harus diisi berurutan' per part, off by default, kept when the
    step is left and shown again, and dropped for parts of other templates."""
    import wizard as wz

    lib, scans, tmp = setup
    errors = []
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: errors.append(a[1]))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.exam.in_order = ["Hilang"]  # part of an earlier template
        w.next()
        step = w.steps[1]
        boxes = [b for b in iter_widgets(step.keys_frame)
                 if isinstance(b, wz.ttk.Checkbutton) and "berurutan" in str(b.cget("text"))]  # fmt: skip
        assert len(boxes) == len(step.texts) and not any(v.get() for v in step.in_order.values())
        first = next(iter(step.texts))
        step.texts[first].insert("1.0", "ABCDE")
        step.on_key_edit(first)
        boxes[0].invoke()
        assert step.in_order[first].get() and w.exam.in_order == {ujian.ALL_CODES: ["Hilang", first]}
        w.next()
        assert w.current == 2, errors
        assert w.exam.in_order == {ujian.ALL_CODES: [first]}
        w.back()
        assert step.in_order[first].get()
        boxes = [b for b in iter_widgets(step.keys_frame)
                 if isinstance(b, wz.ttk.Checkbutton) and "berurutan" in str(b.cget("text"))]  # fmt: skip
        boxes[0].invoke()
        assert w.exam.in_order == {ujian.ALL_CODES: []}
    finally:
        w.destroy()


def test_each_template_keeps_its_own_keys(setup, monkeypatch):
    """Keys typed in step 2 stay with the template chosen in step 1: another
    template starts empty, and going back to the first brings its keys back."""
    import shutil

    import wizard as wz

    lib, scans, tmp = setup
    shutil.copytree(lib / "LJK Sintetis", lib / "LJK Kedua")
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        first = w.steps[0]
        paths = [str(t.path) for t in first.templates]
        first.tree.selection_set("0")
        w.next()
        step = w.steps[1]
        part = next(iter(step.texts))
        step.texts[part].insert("1.0", "ABCDE")
        step.on_key_edit(part)
        step.weights[part].set("1-2:2")
        step.in_order[part].set(True)
        step.on_in_order(part)
        w.back()
        first.tree.selection_set("1")
        w.next()
        assert w.exam.template == paths[1]
        assert step.texts[part].get("1.0", "end").strip() == "" and not step.in_order[part].get()
        assert step.weights[part].get() == "" and w.exam.keys == {}
        w.back()
        first.tree.selection_set("0")
        w.next()
        assert w.exam.template == paths[0]
        assert step.texts[part].get("1.0", "end").strip() == "ABCDE"
        assert step.weights[part].get() == "1-2:2" and step.in_order[part].get()
        w.next()
        saved = ujian.Exam.load(tmp / "sesi.json")
        assert saved.keys[ujian.ALL_CODES] == {part: "ABCDE"} and saved.in_order == {ujian.ALL_CODES: [part]}
    finally:
        w.destroy()


def test_next_exam_after_restart_is_not_exported_yet(setup, monkeypatch):
    """Exported, closed, started again for the next exam (folders not
    remembered): its rekap is "not exported yet" and saving it must not
    suggest the previous exam's rekap file."""
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: False)
    first = tmp / "rekap_pertama.xlsx"
    ujian.Exam(scans_dir=str(scans), output_dir=str(tmp / "hasil1")).save(tmp / "sesi.json")
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        run_to_results(w)
        w.next()
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(first))
        assert w.steps[5].export(ask_open=False)
        w.close()
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass

    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        assert w.exam.exported_file == "" and w.exam.output_dir == ""
        pump(w)
        w.next()
        w.next()
        w.steps[2].folder.set(str(scans))
        w.next()
        w.steps[3].folder.set(str(tmp / "hasil2"))
        w.next()
        step = w.steps[4]
        step.run_btn.invoke()
        deadline = time.time() + 300
        while step.proc is not None:
            w.update()
            time.sleep(0.05)
            assert time.time() < deadline
        w.next()
        assert w.steps[5].export_info.get() == "Belum diekspor"
        asked = []
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: asked.append(k) or "")
        w.steps[5].export(ask_open=False)
        assert Path(asked[0]["initialdir"]) == tmp / "hasil2" and asked[0]["initialfile"] != first.name
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_key_step_mode_switch_keeps_keys_and_in_order_per_code(setup, monkeypatch):
    """Switching to 'Kunci berbeda per kode naskah' used to wipe the one key
    typed before; 'Harus diisi berurutan' did not follow the exam code."""
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    codes = iter(["1", "2"])
    monkeypatch.setattr(wz.simpledialog, "askstring", lambda *a, **k: next(codes))
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.next()
        step = w.steps[1]
        part = next(iter(step.texts))

        def typed():
            return step.texts[part].get("1.0", "end").replace(" ", "").strip()

        def type_key(key):
            step.texts[part].delete("1.0", "end")
            step.texts[part].insert("1.0", key)
            step.on_key_edit(part)

        type_key("ABCDE")
        step.mode.set("code")
        step.on_mode()
        step.add_code()
        type_key("BBBBB")
        step.in_order[part].set(True)
        step.on_in_order(part)
        step.add_code()
        assert step.current_code == "2" and typed() == "" and not step.in_order[part].get()
        type_key("CCCCC")
        step.codes.selection_clear(0, "end")
        step.codes.selection_set(0)
        step.show_code()
        assert typed() == "BBBBB" and step.in_order[part].get()
        step.mode.set("single")
        step.on_mode()
        assert typed() == "ABCDE" and not step.in_order[part].get()
        step.mode.set("code")
        step.on_mode()
        assert list(step.codes.get(0, "end")) == ["1", "2"] and typed() == "BBBBB"
        w.next()  # leave in per-code mode: the one key is kept aside
        assert w.current == 2
        assert w.exam.code_column and ujian.active_codes(w.exam) == ["1", "2"]
        assert w.exam.keys[ujian.ALL_CODES] == {part: "ABCDE"}
        assert w.exam.in_order == {"1": [part]}
        w.back()
        step.mode.set("single")
        step.on_mode()
        assert typed() == "ABCDE"
        w.next()
        assert w.exam.code_column == "" and w.exam.keys["1"] == {part: "BBBBB"}
    finally:
        w.destroy()


def test_key_step_save_and_open_keys_file(setup, monkeypatch):
    """'Simpan kunci...' writes the keys (per exam code too) to a file that
    'Buka kunci...' brings back, e.g. for a new exam on the same LJK."""
    import wizard as wz

    lib, scans, tmp = setup
    shown = []
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: pytest.fail(f"error: {a}"))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(wz.messagebox, "showinfo", lambda *a, **k: shown.append(a[1]))
    monkeypatch.setattr(wz.messagebox, "showwarning", lambda *a, **k: shown.append(a[1]))
    monkeypatch.setattr(wz.simpledialog, "askstring", lambda *a, **k: "7")
    path = tmp / f"kunci{ujian.KEYS_EXT}"
    monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(path))
    monkeypatch.setattr(wz.filedialog, "askopenfilename", lambda **k: str(path))
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.next()
        step = w.steps[1]
        buttons = {str(b.cget("text")): b for b in iter_widgets(step) if isinstance(b, wz.ttk.Button)}
        buttons["Simpan kunci..."].invoke()
        assert not path.exists() and shown == ["Belum ada kunci jawaban yang bisa disimpan."]
        parts = list(step.texts)
        step.mode.set("code")
        step.on_mode()
        step.add_code()
        for p in parts:
            step.texts[p].insert("1.0", "D")
            step.on_key_edit(p)
        step.in_order[parts[0]].set(True)
        step.on_in_order(parts[0])
        step.weights[parts[0]].set("1:3")
        step.correct.set("2")
        buttons["Simpan kunci..."].invoke()
        assert path.exists()

        # new exam: everything gone, then the file brings it back
        w.exam.keys, w.exam.in_order, w.exam.weights, w.exam.code_column = {}, {}, {}, ""
        step.enter()
        assert step.mode.get() == "single" and step.texts[parts[0]].get("1.0", "end").strip() == ""
        shown.clear()
        buttons["Buka kunci..."].invoke()
        assert shown == []
        assert step.mode.get() == "code" and list(step.codes.get(0, "end")) == ["7"]
        assert step.texts[parts[0]].get("1.0", "end").strip() == "D" and step.in_order[parts[0]].get()
        assert step.weights[parts[0]].get() == "1:3" and step.correct.get() == "2"
        w.next()
        assert w.current == 2 and w.exam.keys == {"7": {p: "D" for p in parts}} and w.exam.correct == 2
    finally:
        w.destroy()
