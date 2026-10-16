"""The wizard, driven step by step like a user, on the Medion LJK: template
from the library, keys per part, scan folder + student list, output folder,
real processing in the worker, a manual correction and the Excel export.

Needs Tk and a display (skipped otherwise); headless Linux: xvfb-run.
OMR_SCREENSHOT_DIR=<dir> keeps a screenshot of every step.
"""

import os
import subprocess
import sys
import time
from pathlib import Path

import cv2
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(Path(__file__).parent)]

tk = pytest.importorskip("tkinter", exc_type=ImportError)

import synth_ljk as synth  # noqa: E402
import template_builder as tb  # noqa: E402
import test_medion_ljk as medion  # noqa: E402
import ujian  # noqa: E402

from testdata import needs  # noqa: E402

pytestmark = [needs("ljk_medion_terisi.pdf")]


def iter_widgets(w):
    for c in w.winfo_children():
        yield c
        yield from iter_widgets(c)


def pump(w, n=10):
    for _ in range(n):
        w.update()


def shot(w, name):
    folder = os.environ.get("OMR_SCREENSHOT_DIR")
    if folder:
        w.geometry("1200x780+0+0")
        pump(w, 20)
        subprocess.run(["import", "-window", "root", str(Path(folder, f"wizard_{name}.png"))], check=False)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    try:
        tk.Tk().destroy()
    except tk.TclError as e:
        pytest.skip(f"no display: {e}")
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    sheet = tb.prepare_sheet(tb.load_image(medion.SCAN))
    bubbles, blocks = tb.detect_blocks(sheet)
    blocks = medion.label_like_a_teacher(bubbles, blocks)
    template, _ = tb.build_template(blocks, sheet)
    lib = tmp_path / "lib"
    ujian.save_to_library("LJK Medion", sheet, template, blocks, lib)

    scans = tmp_path / "scan"
    scans.mkdir()
    (scans / "siswa.pdf").write_bytes(medion.SCAN.read_bytes())
    img = synth.distort(tb.load_image(medion.SCAN), angle=1.0, dx=20, jpeg=85)
    cv2.imencode(".jpg", img)[1].tofile(str(scans / "siswa_miring.jpg"))

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["No", "Nama", "Kelas"])
    wb.active.append([1, medion.EXPECTED.get("Nama", "").title(), "XII TE"])
    wb.active.append([2, "Siswa Absen", "XII TE"])
    wb.save(tmp_path / "siswa.xlsx")
    return lib, scans, tmp_path


def test_wizard_full_flow(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    errors = []
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: errors.append(a))
    monkeypatch.setattr(wz.messagebox, "showinfo", lambda *a, **k: errors.append(a))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: False)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        # 1. template
        step = w.steps[0]
        assert len(step.templates) == 1 and step.tree.selection() == ("0",)
        assert step.preview.img is not None
        shot(w, "1_template")
        w.next()
        assert w.current == 1 and not errors

        # 2. keys per part (the pencilled answers), KKM
        step = w.steps[1]
        assert set(step.texts) == {"I", "II", "III", "Motivasi", "TYPE"}
        for part, key in medion.ANSWERS.items():
            step.texts[part].insert("1.0", key)
            step.on_key_edit(part)
        assert step.counters["I"].get().startswith("20 / 130")
        step.weights["I"].set("1-10:2")  # q1-q10 of part I worth 2 points
        pump(w)
        assert "skor maksimal 30" in step.counters["I"].get()
        step.kkm.set("60")
        shot(w, "2_kunci")
        w.next()
        assert w.current == 2, errors

        # 3. scans + student list (no NIS on this sheet: matched by name)
        step = w.steps[2]
        step.folder.set(str(scans))
        pump(w)
        assert step.count.get().startswith("2 file")
        step.students.set(str(tmp / "siswa.xlsx"))
        pump(w)
        assert step.col_vars["nama"].get() == "Nama" and "2 baris" in step.student_info.get()
        shot(w, "3_folder")
        w.next()
        assert w.current == 3, errors

        # 4. output folder (default next to the scans)
        step = w.steps[3]
        assert step.folder.get().endswith("scan_hasil")
        step.name.set("Tes Medion 2026")
        shot(w, "4_hasil")
        w.next()
        assert w.current == 4, errors

        # 5. processing waits for the click
        step = w.steps[4]
        pump(w, 20)
        assert step.proc is None and not step.results
        assert "Klik" in step.summary.get()
        step.run_btn.invoke()
        deadline = time.time() + 300
        while (step.proc is None and not step.results) or step.proc is not None:
            w.update()
            time.sleep(0.05)
            assert time.time() < deadline, "processing did not finish"
        pump(w, 20)
        assert not errors, errors
        assert len(step.results) == 2
        for r in step.results:
            assert r.info["Nama"].strip() == medion.EXPECTED.get("Nama") and r.info["KodeDuduk"] == "A1", r.file_id
            assert r.grade == 100 and r.student is not None and r.kelas == "XII TE", r.flags
            assert r.parts["I"]["skor"] == 30  # 10 x 2 + 10 x 1
        assert step.preview.img is not None
        assert step.preview.winfo_width() > 250, "preview pane squeezed out"
        shot(w, "5_proses")

        # manual correction: one answer changed -> part I loses a point
        step.tree.selection_set("0")
        pump(w)
        step.correct()
        pump(w)
        dlg = next(c for c in step.winfo_children() if isinstance(c, wz.CorrectionDialog))
        dlg.geometry("900x500")  # short window: save button must stay visible
        pump(w)
        save_btn = next(b for b in iter_widgets(dlg) if isinstance(b, wz.ttk.Button) and b.cget("text") == "Simpan koreksi")
        assert save_btn.winfo_ismapped() and save_btn.winfo_rooty() + save_btn.winfo_height() <= dlg.winfo_rooty() + dlg.winfo_height()
        text = dlg.texts["I"].get("1.0", "end")
        dlg.texts["I"].delete("1.0", "end")
        dlg.texts["I"].insert("1.0", "B" + text[1:])  # q1 A -> B
        dlg.save()
        pump(w)
        fixed = next(r for r in step.results if r.corrected)
        assert fixed.parts["I"]["benar"] == 19 and fixed.grade < 100
        assert fixed.parts["I"]["skor"] == 28  # q1 was worth 2
        assert w.exam.corrections[fixed.file_id] == {"I_1": "B"}
        shot(w, "5b_koreksi")
        w.next()
        assert w.current == 5, errors

        # 6. export
        step = w.steps[5]
        assert "2 LJK" in step.stats.get() and "tidak ikut: 1" in step.stats.get()
        assert len(step.tree.get_children()) == sum(map(len, medion.ANSWERS.values()))  # keyed only
        out = tmp / "rekap.xlsx"
        monkeypatch.setattr(wz.filedialog, "asksaveasfilename", lambda **k: str(out))
        step.export()
        shot(w, "6_rekap")
        from openpyxl import load_workbook

        wb = load_workbook(out)
        header = [c.value for c in wb["Rekap Nilai"][5]]
        assert "Benar I" in header and "Nilai Motivasi" in header and "KodeDuduk" in header
        assert [c.value for c in wb["Tidak Ikut Ujian"][2]][2] == "Siswa Absen"

        # settings survive a restart
        w.on_close()
        again = ujian.Exam.load(tmp / "sesi.json")
        assert again.scans_dir == str(scans) and again.keys[ujian.ALL_CODES]["III"] == "ABACDCBB"
    finally:
        try:
            w.destroy()
        except tk.TclError:
            pass


def test_wizard_blocks_bad_input(setup, monkeypatch):
    import wizard as wz

    lib, scans, tmp = setup
    errors = []
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: errors.append(a[1]))
    monkeypatch.setattr(wz.messagebox, "askyesno", lambda *a, **k: True)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        pump(w)
        w.next()
        step = w.steps[1]
        step.texts["TYPE"].insert("1.0", "DD")  # 2 answers for 1 question
        step.on_key_edit("TYPE")
        w.next()
        assert w.current == 1 and "kunci berisi 2 jawaban, padahal bagian ini hanya 1 soal" in errors[-1]
        step.weights["TYPE"].set("1-5:2")  # TYPE has a single question
        w.next()
        assert w.current == 1 and "hanya sampai 1" in errors[-1]
        step.weights["TYPE"].set("")
        step.texts["TYPE"].delete("1.0", "end")
        step.on_key_edit("TYPE")
        w.next()  # no key at all: confirmed to just read the sheets
        assert w.current == 2
        w.steps[2].folder.set(str(tmp / "tidak-ada"))
        w.next()
        assert w.current == 2
        w.steps[2].folder.set(str(scans))
        w.next()
        w.steps[3].folder.set(str(scans / "hasil"))  # inside the scans: refused
        w.next()
        assert w.current == 3 and "terpisah" in errors[-1]
    finally:
        w.destroy()


def test_key_step_shows_parts_and_renames_them(setup, monkeypatch):
    """Step 2 shows where each part is on the sheet and renames a part in
    place (the reported confusion: auto names IV/V vs TYPE/III)."""
    import wizard as wz

    lib, scans, tmp = setup
    monkeypatch.setattr(wz.messagebox, "showerror", lambda *a, **k: None)
    w = wz.run_wizard(mainloop=False, library=lib, settings=tmp / "sesi.json")
    try:
        w.geometry("1000x620+0+0")  # small screen
        pump(w)
        w.next()
        step = w.steps[1]
        assert len(step.thumbs) == 5  # a crop of the sheet beside every key
        step.texts["TYPE"].insert("1.0", "D")
        step.on_key_edit("TYPE")
        step.weights["TYPE"].set("1:3")
        monkeypatch.setattr(wz.simpledialog, "askstring", lambda *a, **k: "Tipe")
        step.rename("TYPE")
        pump(w)
        assert "Tipe" in step.texts and "TYPE" not in step.texts
        assert w.template.sections["Tipe"] == ["Tipe_1"]
        assert w.exam.keys[ujian.ALL_CODES]["Tipe"] == "D" and w.exam.weights["Tipe"] == "1:3"
        assert step.texts["Tipe"].get("1.0", "end").strip() == "D"
        assert ujian.read_library_template(w.template.path).sections["Tipe"] == ["Tipe_1"]
        # the Next button stays on screen
        assert w.next_btn.winfo_rooty() + w.next_btn.winfo_height() <= w.winfo_rooty() + w.winfo_height()
        shot(w, "2b_kunci_gambar")
    finally:
        w.destroy()
